"""Export test design fields without turning planned cases into test results."""

import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Font

FIELDS = {
    "id": "ID",
    "title": "Title",
    "status": "Status",
    "precondition": "Precondition",
    "inputs": "Inputs",
    "expected_outputs": "Expected Outputs",
    "pass_criterion": "Pass Criterion",
    "req_id": "Requirement ID",
    "test_level": "Test Level",
    "environment": "Environment",
    "automation": "Automation",
    "priority": "Priority",
}


def cell_text(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


def main(input_json, output_xlsx):
    with open(input_json, encoding="utf-8") as stream:
        data = json.load(stream)
    cases = data.get("test_cases", [])
    if not isinstance(cases, list) or any(not isinstance(tc, dict) for tc in cases):
        raise ValueError("test_cases must be a list of objects")
    wb = Workbook()
    title = wb.active
    title.title = "Title"
    title.append(["Test Case Catalog"])
    for key, value in data.get("item", {}).items():
        title.append([key, cell_text(value)])
    ws = wb.create_sheet("Test Case Inventory")
    ws.append(list(FIELDS.values()))
    for tc in cases:
        ws.append([cell_text(tc.get(key, "")) for key in FIELDS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    references = wb.create_sheet("References")
    references.append(["Source", "Value"])
    for key, value in data.get("provenance", {}).items():
        references.append([key, cell_text(value)])
    wb.save(output_xlsx)
    print(f"Test Case Catalog generated: {output_xlsx} ({len(cases)} planned cases)")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("Usage: generate_tc_catalog.py <input.json> <output.xlsx>")
    main(sys.argv[1], sys.argv[2])
