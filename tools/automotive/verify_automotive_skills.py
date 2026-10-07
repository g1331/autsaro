"""Verify optional automotive document skills."""

import ast
import copy
import importlib.util
import io
import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from contextlib import contextmanager
from pathlib import Path
from types import SimpleNamespace

from install_automotive_skills import MANIFEST, ROOT, extract, validate
from openpyxl import load_workbook

SKILLS = ROOT / ".agents/skills"
ITEM = {
    "name": "Synthetic fixture",
    "project": "Installation smoke test",
    "doc_id": "FIXTURE-001",
    "revision": "draft",
    "author": "Fixture",
    "approver": "Unapproved",
    "date": "2026-09-28",
}
CASE = {
    "id": "TC-1",
    "title": "Synthetic success",
    "status": "Not executed",
    "precondition": "Fresh temporary environment",
    "inputs": "uint32 42",
    "expected_outputs": "42",
    "pass_criterion": "exact equality",
    "req_id": "REQ-1",
    "test_level": "integration",
    "environment": "synthetic host",
    "automation": "planned",
    "priority": "must",
}


def fixtures():
    return {
        "traceability-matrix": {
            "stakeholder_needs": [{"id": "NEED-1", "title": "Fixture need"}],
            "system_requirements": [
                {"id": "REQ-1", "title": "Fixture requirement", "traces_to": ["NEED-1"]}
            ],
            "test_cases": [
                {
                    "id": "TC-1",
                    "title": "Fixture case",
                    "verifies": ["REQ-1"],
                    "result": "Not executed",
                }
            ],
        },
        "test-case-catalog": {"test_cases": [CASE]},
        "verification-plan": {
            "scope_description": "Synthetic host scope",
            "requirements": [{"id": "REQ-1", "title": "Fixture requirement"}],
            "verification_methods": [
                {"method": "test", "applicability": "REQ-1", "level": "integration"}
            ],
            "environments": {"sil": {"desc": "Synthetic host", "tools": "Fixture"}},
            "test_levels": [{"level": "integration", "description": "Fixture level"}],
            "acceptance_criteria": [
                {
                    "criterion": "Equality",
                    "measurement": "value",
                    "pass_threshold": "42",
                }
            ],
        },
        "uds-services": {
            "ecu_identity": {"physical_address": "0x700", "logical_address": "fixture"},
            "session_definitions": [{"type": "default", "id": "0x01"}],
            "service_inventory": [
                {
                    "service_id": "0x22",
                    "name": "ReadDataByIdentifier",
                    "session_restrictions": "default",
                }
            ],
            "did_catalog": [{"id": "0xF100", "name": "Fixture", "length_bytes": 4}],
            "negative_response_codes": [{"code": "0x31", "name": "RequestOutOfRange"}],
            "timing_parameters": {"P2_ms": 50, "P2_star_ms": 5000, "S3_ms": 5000},
        },
        "dtc-catalog": {
            "ecu": ITEM,
            "dtcs": [
                {
                    "id": "DTC-1",
                    "code": "0x000100",
                    "name": "Synthetic fault",
                    "description": "Fixture only",
                }
            ],
        },
        "dem-config": {
            "events": [
                {"event_id": "E-1", "dtc": "0x000100", "operation_cycle": "Fixture"}
            ],
            "debounce_config": [{"event_id": "E-1", "confirmed": 1, "healed": 0}],
            "memory_config": {"primary_size_bytes": 1024},
        },
        "arxml-system": {
            "ecu_instances": [{"name": "ECU1", "id": "ECU-1"}],
            "signal_catalog": [{"name": "Value", "id": "SIG-1", "length": 32}],
            "pdu_catalog": [{"name": "Pdu1", "id": "PDU-1"}],
            "frame_catalog": [{"name": "Frame1", "id": "FRAME-1"}],
        },
        "autosar-swc": {
            "swc_type": {"name": "FixtureSWC", "category": "Application"},
            "ports": [
                {
                    "name": "Value",
                    "direction": "Provided",
                    "category": "SR",
                    "interface_type": "SenderReceiver",
                    "interface_name": "ValueInterface",
                    "data_element": "Value",
                }
            ],
            "runnables": [{"name": "Run", "id": "RUN-1"}],
            "events": [
                {"name": "Tick", "id": "EV-1", "type": "TimingEvent", "period_ms": 10}
            ],
        },
        "autosar-composition": {
            "composition": {"name": "FixtureComposition"},
            "swc_instances": [{"name": "App", "swc_type": "FixtureSWC"}],
            "boundary_ports": [
                {
                    "name": "Value",
                    "direction": "Provided",
                    "interface_name": "ValueInterface",
                }
            ],
        },
        "autosar-rte-mapping": {
            "ecu": {"name": "ECU1"},
            "data_mappings": [{"swc": "App", "port": "Value", "com_signal": "Signal1"}],
            "runnable_to_task": [{"runnable": "Run", "task": "Task1"}],
            "schedulable_tasks": [{"name": "Task1", "priority": 1, "period_ms": 10}],
        },
        "autosar-bsw-config": {
            "ecu_target": {"microcontroller": "Synthetic target", "vendor": "Fixture"},
            "mcal_modules": [
                {"id": "CAN", "name": "Can", "configured_for_target": True}
            ],
            "module_parameters": [
                {"module_id": "CAN", "parameter": "CanControllerId", "value": 0}
            ],
            "os_tasks": [
                {"id": "T-1", "name": "Task1", "priority": 1, "period_ms": 10}
            ],
        },
        "secure-coding-guidelines": {
            "project": ITEM,
            "code_base": "Synthetic C99 fixture",
            "cal": "Not assessed",
            "language": "C",
            "applicable_standards": ["CERT C"],
            "static_analysis_tools": [],
        },
    }


def execute(script, *args):
    result = subprocess.run(
        [sys.executable, str(script), *map(str, args)],
        cwd=ROOT,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=60,
        check=False,
    )
    if result.returncode:
        raise AssertionError(
            f"{script.parent.parent.name}: {result.stdout}\n{result.stderr}"
        )
    return result


def generate(base, data, directory):
    data = {
        "item": ITEM,
        "provenance": {"status": "Design fixture; NOT EXECUTED"},
        **data,
    }
    input_path = directory / f"{base}.json"
    input_path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    builder = next((SKILLS / f"{base}-builder/scripts").glob("generate_*.py"))
    output = directory / f"{base}.xlsx"
    execute(builder, input_path, output)
    reviewer = next(
        (SKILLS / f"{base}-checklist-reviewer/scripts").glob("generate_*.py")
    )
    report = directory / f"{base}-review.xlsx"
    execute(reviewer, output, report)
    return output, report


@contextmanager
def module(base, filename):
    """Isolate same-named upstream modules from other skill pairs."""
    path = SKILLS / f"{base}-checklist-reviewer/scripts"
    saved = sys.path[:]
    sys.path.insert(0, str(path))
    spec = importlib.util.spec_from_file_location("skill_under_test", path / filename)
    loaded = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = loaded
    try:
        spec.loader.exec_module(loaded)
        yield loaded
    finally:
        sys.path[:] = saved
        del sys.modules[spec.name]


class InstalledSkillsTests(unittest.TestCase):
    def test_inventory_and_syntax(self):
        lock = json.loads(MANIFEST.read_text(encoding="utf-8"))
        self.assertEqual(len(lock["skills"]), 24)
        validate(lock)
        count = 0
        for name in lock["skills"]:
            root = SKILLS / name
            entry = (root / "SKILL.md").read_text(encoding="utf-8")
            self.assertTrue(entry.startswith("---\n"))
            self.assertIn(f"name: {name}\n", entry)
            self.assertIn("description:", entry)
            self.assertTrue((root / "LICENSE").is_file())
            for path in root.rglob("*.py"):
                ast.parse(path.read_bytes(), filename=str(path))
                count += 1
        print(f"Validated 24 skill entries and {count} Python files")

    def test_all_builder_reviewer_pairs(self):
        with tempfile.TemporaryDirectory() as directory:
            for base, data in fixtures().items():
                with self.subTest(skill=base):
                    output, report = generate(base, data, Path(directory))
                    for path in (output, report):
                        wb = load_workbook(path)
                        try:
                            self.assertIn("Repository Scope", wb.sheetnames)
                            self.assertGreater(len(wb.sheetnames), 2)
                            values = [
                                v
                                for ws in wb
                                for row in ws.iter_rows(values_only=True)
                                for v in row
                                if v is not None
                            ]
                            self.assertTrue(values)
                            if path == output:
                                expected = {
                                    "autosar-composition": "FixtureComposition",
                                    "autosar-swc": "FixtureSWC",
                                }.get(base, "Synthetic fixture")
                                self.assertIn(expected, str(values))
                                markers = {
                                    "traceability-matrix": "REQ-1",
                                    "test-case-catalog": "uint32 42",
                                    "verification-plan": "Synthetic host scope",
                                    "uds-services": "0xF100",
                                    "dtc-catalog": "0x000100",
                                    "dem-config": "E-1",
                                    "arxml-system": "SIG-1",
                                    "autosar-swc": "ValueInterface",
                                    "autosar-composition": "App",
                                    "autosar-rte-mapping": "Signal1",
                                    "autosar-bsw-config": "CanControllerId",
                                    "secure-coding-guidelines": "Synthetic C99 fixture",
                                }
                                self.assertIn(markers[base], values)
                                if base == "traceability-matrix":
                                    self.assertIn("Not executed", values)
                                self.assertTrue(
                                    all(len(name) <= 31 for name in wb.sheetnames)
                                )
                        finally:
                            wb.close()
                    if base == "traceability-matrix":
                        wb = load_workbook(report)
                        try:
                            self.assertEqual(
                                wb["Execution Evidence"]["B2"].value, "Not executed"
                            )
                        finally:
                            wb.close()
                    print(f"PASS builder/reviewer: {base}")

    def test_test_catalog_reads_design_and_rejects_gaps(self):
        with tempfile.TemporaryDirectory() as directory:
            data = fixtures()["test-case-catalog"]
            data = copy.deepcopy(data)
            data["test_cases"] += [dict(CASE, expected_outputs="")]
            output, _ = generate("test-case-catalog", data, Path(directory))
            with module("test-case-catalog", "tc_probe.py") as probe:
                result = probe.probe_tc_catalog(output)
            self.assertEqual(result.test_cases[0].inputs, "uint32 42")
            with module("test-case-catalog", "check_definitions.py") as checks:
                ratings = checks.evaluate_tc_catalog(result)
                self.assertEqual(ratings["TC001"]["rating"], "NO")
                self.assertEqual(ratings["TC005"]["rating"], "NO")
                empty = checks.evaluate_tc_catalog(SimpleNamespace(test_cases=[]))
                self.assertTrue(all(v["rating"] == "NA" for v in empty.values()))

    def test_trace_orphan_and_missing_metric(self):
        with tempfile.TemporaryDirectory() as directory:
            data = copy.deepcopy(fixtures()["traceability-matrix"])
            data["system_requirements"].append(
                {"id": "REQ-ORPHAN", "title": "Unverified"}
            )
            output, _ = generate("traceability-matrix", data, Path(directory))
            with module("traceability-matrix", "trace_probe.py") as probe:
                result = probe.probe_trace_matrix(output)
            with module("traceability-matrix", "check_definitions.py") as checks:
                ratings = checks.evaluate_trace_matrix(result)
                self.assertEqual(ratings["TM005"]["rating"], "NO")
                result.metrics = {}
                self.assertEqual(
                    checks.evaluate_trace_matrix(result)["TM007"]["rating"], "NA"
                )

    def test_bsw_parameter_conflict(self):
        with tempfile.TemporaryDirectory() as directory:
            data = copy.deepcopy(fixtures()["autosar-bsw-config"])
            data["module_parameters"] *= 2
            output, _ = generate("autosar-bsw-config", data, Path(directory))
            with module("autosar-bsw-config", "bsw_probe.py") as probe:
                result = probe.probe(output)
            with module("autosar-bsw-config", "check_definitions.py") as checks:
                verdicts = [check.execute(result) for check in checks.CHECKS]
            self.assertTrue(
                any(v.rating == "NO" and "Duplicate" in v.finding for v in verdicts)
            )

    def test_unassessed_plan_and_rte_analysis(self):
        with module("verification-plan", "check_definitions.py") as checks:
            ratings = checks.evaluate_vplan(
                SimpleNamespace(
                    scope="",
                    methods=[],
                    environments={},
                    test_levels=[],
                    acceptance_criteria=[],
                )
            )
            self.assertEqual(ratings["VP025"]["rating"], "NA")
        with module("autosar-rte-mapping", "check_definitions.py") as checks:
            result = checks.check_priority_inversion(
                SimpleNamespace(schedulable_tasks=[{"priority": 1}, {"priority": 2}])
            )
            self.assertEqual(result.rating, "NA")

    def test_installer_detects_local_changes(self):
        from unittest.mock import patch

        import install_automotive_skills as installer

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            target = root / ".agents/skills/demo"
            target.mkdir(parents=True)
            source = target / "SKILL.md"
            source.write_bytes(b"original\n")
            lock = {
                "file_hash_mode": "lf-normalized-text",
                "skills": {"demo": {"files": installer.inventory(target)}},
            }
            with patch.object(installer, "ROOT", root):
                validate(lock)
                source.write_bytes(b"original\r\n")
                validate(lock)
                source.write_text("local modification", encoding="utf-8")
                with self.assertRaisesRegex(ValueError, "refusing overwrite"):
                    validate(lock)
                self.assertEqual(
                    source.read_text(encoding="utf-8"), "local modification"
                )

    def test_unsafe_archive_never_writes(self):
        with tempfile.TemporaryDirectory() as directory:
            for member in (
                "demo/../escape",
                "demo/C:/escape",
                "demo\\escape",
                "/demo/escape",
            ):
                data = io.BytesIO()
                with zipfile.ZipFile(data, "w") as archive:
                    archive.writestr("demo/SKILL.md", "valid")
                    info = zipfile.ZipInfo("placeholder")
                    info.filename = member
                    archive.writestr(info, "unsafe")
                with self.subTest(member=member):
                    with self.assertRaises(ValueError):
                        extract(data.getvalue(), "demo", Path(directory) / "out")
                    self.assertFalse((Path(directory) / "out").exists())


def demo():
    import hashlib

    story_path = (
        ROOT
        / "_bmad-output/implementation-artifacts/3-3-verify-delivered-host-behavior-without-the-workbench.md"
    )
    story = story_path.read_bytes()
    commit = subprocess.check_output(
        ["git", "rev-parse", "HEAD"], cwd=ROOT, text=True
    ).strip()
    provenance = {
        "source": story_path.relative_to(ROOT).as_posix(),
        "git_commit": commit,
        "source_sha256": hashlib.sha256(story).hexdigest(),
        "status": "设计演练，尚未执行运行用例",
    }
    designs = [
        (
            "TC-CAN",
            "双 ECU CAN 双向信号",
            "AC-1",
            "固定参考 CAN 帧及信号输入",
            "随包独立参考数据中的精确 CAN 帧",
        ),
        (
            "TC-UDS",
            "合法物理诊断请求",
            "AC-2",
            "声明支持的物理诊断请求",
            "随包独立请求/响应预期",
        ),
        (
            "TC-RECOVERY",
            "拒绝或故障恢复",
            "AC-2",
            "适用的不支持请求或传输故障向量",
            "明确拒绝或恢复后的合法响应",
        ),
        (
            "TC-PACKAGE",
            "篡改文件失败",
            "AC-3",
            "删除或篡改交付文件",
            "验证入口非零退出并保存定位记录",
        ),
    ]
    cases = [
        dict(
            CASE,
            id=tid,
            title=title,
            req_id=req,
            inputs=inputs,
            expected_outputs=expected,
            status="Not executed",
            pass_criterion="按固定独立参考数据逐项比较；精确字节须在执行前补齐",
            precondition="Epic 3.3 固定参考包，干净临时目录，声明的 Windows/PowerShell/GCC",
            environment="Windows host; design only",
        )
        for tid, title, req, inputs, expected in designs
    ]
    trace = {
        "provenance": provenance,
        "stakeholder_needs": [
            {"id": "NEED-HANDOFF", "title": "接收者能离线复验参考主机工程"}
        ],
        "system_requirements": [
            {"id": req, "title": title, "traces_to": ["NEED-HANDOFF"]}
            for req, title in [
                ("AC-1", "双 ECU CAN 独立比对"),
                ("AC-2", "诊断合法及拒绝/恢复"),
                ("AC-3", "失败关闭及记录"),
                ("AC-4", "支持边界声明（本演练未覆盖）"),
            ]
        ],
        "test_cases": [
            {
                "id": tc["id"],
                "title": tc["title"],
                "verifies": [tc["req_id"]],
                "result": "NOT EXECUTED",
            }
            for tc in cases
        ],
    }
    with tempfile.TemporaryDirectory(prefix="autosar-epic3-design-") as directory:
        catalog, _ = generate(
            "test-case-catalog",
            {"test_cases": cases, "provenance": provenance},
            Path(directory),
        )
        matrix, _ = generate("traceability-matrix", trace, Path(directory))
        with module("test-case-catalog", "tc_probe.py") as probe:
            assert len(probe.probe_tc_catalog(catalog).test_cases) == 4
        with module("traceability-matrix", "trace_probe.py") as probe:
            result = probe.probe_trace_matrix(matrix)
        assert any(row["id"] == "AC-4" for row in result.orphan_reqs)
    print(
        "Epic 3.3 design demo: 4 unexecuted cases; uncovered AC-4 correctly retained; no task/capability changes"
    )


if __name__ == "__main__":
    if sys.argv[1:] == ["--demo"]:
        demo()
    else:
        unittest.main(verbosity=2)
