"""Export test design fields without turning planned cases into test results."""

import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Font

FIELDS = {
    "id": "ID",
    "title": "Title",
    "status": "Status",
    "precondition": "Precondition",
    "inputs": "Inputs",
    "expected_outputs": "Expected Outputs",
    "pass_criterion": "Pass Criterion",
    "req_id": "Requirement ID",
    "test_level": "Test Level",
    "environment": "Environment",
    "automation": "Automation",
    "priority": "Priority",
}


def cell_text(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


def main(input_json, output_xlsx):
    with open(input_json, encoding="utf-8") as stream:
        data = json.load(stream)
    cases = data.get("test_cases", [])
    if not isinstance(cases, list) or any(not isinstance(tc, dict) for tc in cases):
        raise ValueError("test_cases must be a list of objects")
    wb = Workbook()
    title = wb.active
    title.title = "Title"
    title.append(["Test Case Catalog"])
    for key, value in data.get("item", {}).items():
        title.append([key, cell_text(value)])
    ws = wb.create_sheet("Test Case Inventory")
    ws.append(list(FIELDS.values()))
    for tc in cases:
        ws.append([cell_text(tc.get(key, "")) for key in FIELDS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    references = wb.create_sheet("References")
    references.append(["Source", "Value"])
    for key, value in data.get("provenance", {}).items():
        references.append([key, cell_text(value)])
    _repository_notice(wb)
    wb.save(output_xlsx)
    print(f"Test Case Catalog generated: {output_xlsx} ({len(cases)} planned cases)")




def _repository_notice(wb):
    import json
    from pathlib import Path
    import sys
    for sheet in wb:
        for row in sheet:
            for cell in row:
                if cell.value == "APPROVED":
                    cell.value = "DOCUMENT CHECKS COMPLETE - review required"
                elif cell.value == "CONDITIONAL APPROVAL":
                    cell.value = "DOCUMENT ISSUES - review required"
                elif cell.value in ("Internal review comments incorporated", "Released for architecture review", "Initial release"):
                    cell.value = "TEMPLATE PLACEHOLDER - not a recorded project event"
    if "03_Coding_Rules_Catalog" in wb.sheetnames:
        rules = wb["03_Coding_Rules_Catalog"]
        col = rules.max_column + 1
        rules.cell(1, col, "Repository validation status")
        for row in range(2, rules.max_row + 1):
            rules.cell(row, col, "UNVERIFIED UPSTREAM TEMPLATE - rule ID/text/severity/CWE require authoritative review")
    ws = wb.create_sheet("Repository Scope")
    ws.append(["Field", "Value"])
    ws.append(["Purpose", "Document generation/review only; NOT a runtime test or certification"])
    ws.append(["Evidence", "Separate CONTENT / STRUCTURE / DRAFT; missing evidence remains unassessed"])
    ws.append(["Version", "Upstream AUTOSAR R22-11 templates; repository R24-11 obligations need verification"])
    ws.append(["Thresholds", "Upstream thresholds are advisory, not repository acceptance gates"])
    ws.append(["Templates", "Fixed example/default content requires review; not project facts"])
    if len(sys.argv) > 1 and Path(sys.argv[1]).suffix.lower() == ".json":
        with open(sys.argv[1], encoding="utf-8") as stream:
            provenance = json.load(stream).get("provenance", {})
        for key, value in provenance.items():
            ws.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 100


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("Usage: generate_tc_catalog.py <input.json> <output.xlsx>")
    main(sys.argv[1], sys.argv[2])
