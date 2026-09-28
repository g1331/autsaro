"""Visual dashboard for checklist reviewer skills."""

from __future__ import annotations

from collections import Counter

from openpyxl.chart import PieChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

FONT_NAME = "Calibri"
NAVY = "1F3864"
GREEN_OK = "C6EFCE"
RED_BAD = "F8CBAD"
ORANGE_PARTIAL = "FFD966"
GREY = "D0CECE"

def style_title(ws, row, last_col, text):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(name=FONT_NAME, size=14, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 28

def build_dashboard(wb, results):
    ws = wb.create_sheet("Summary", 1)
    style_title(ws, 1, 8, "Summary Dashboard")

    all_results = []
    for tab_results in results.values():
        for _, result in tab_results:
            all_results.append(result)

    rating_counts = Counter(r.rating for r in all_results)
    total_checks = len(all_results)
    compliant_count = rating_counts.get("FC", 0)
    compliant_pct = (compliant_count / total_checks * 100) if total_checks > 0 else 0

    row = 3
    ws.cell(row=row, column=1, value="Total Checks").font = Font(bold=True)
    ws.cell(row=row, column=2, value=total_checks)
    row += 1
    ws.cell(row=row, column=1, value="Compliant %").font = Font(bold=True)
    ws.cell(row=row, column=2, value=f"{compliant_pct:.0f}%")
    row += 1
    ws.cell(row=row, column=1, value="Major Issues").font = Font(bold=True)
    ws.cell(row=row, column=2, value=rating_counts.get("NO", 0))
