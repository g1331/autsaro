"""Probe AUTOSAR BSW Config xlsx and extract key structures."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

@dataclass
class ProbedBSW:
    title: str
    ecu_target: dict
    mcal_modules: list
    ecu_abstraction: list
    service_layer: list
    complex_drivers: list
    module_parameters: list
    os_tasks: list
    memory_regions: list

def probe(xlsx_path: str) -> ProbedBSW:
    wb = load_workbook(xlsx_path)

    title = ""
    if "Title" in wb.sheetnames:
        ws = wb["Title"]
        title = ws.cell(3, 1).value or "BSW Config"

    ecu_target = {}
    if "ECU-C Top-Level" in wb.sheetnames:
        ws = wb["ECU-C Top-Level"]
        for row in range(3, 10):
            label = ws.cell(row, 1).value
            value = ws.cell(row, 2).value
            if label:
                ecu_target[label.lower().replace(" ", "_")] = value

    mcal_modules = []
    if "MCAL Module Inventory" in wb.sheetnames:
        ws = wb["MCAL Module Inventory"]
        for row in range(3, 100):
            id_val = ws.cell(row, 1).value
            if not id_val:
                break
            mcal_modules.append({
                "id": id_val,
                "name": ws.cell(row, 2).value,
                "configured": ws.cell(row, 3).value,
            })

    ecu_abstraction = []
    if "ECU Abstraction Inventory" in wb.sheetnames:
        ws = wb["ECU Abstraction Inventory"]
        for row in range(3, 100):
            id_val = ws.cell(row, 1).value
            if not id_val:
                break
            ecu_abstraction.append({
                "id": id_val,
                "name": ws.cell(row, 2).value,
                "depends_on": ws.cell(row, 3).value,
            })

    service_layer = []
    if "Service Layer Modules" in wb.sheetnames:
        ws = wb["Service Layer Modules"]
        for row in range(3, 100):
            id_val = ws.cell(row, 1).value
            if not id_val:
                break
            service_layer.append({
                "id": id_val,
                "name": ws.cell(row, 2).value,
                "critical": ws.cell(row, 3).value == "Yes",
            })

    complex_drivers = []
    if "Complex Drivers" in wb.sheetnames:
        ws = wb["Complex Drivers"]
        for row in range(3, 100):
            id_val = ws.cell(row, 1).value
            if not id_val:
                break
            complex_drivers.append({
                "id": id_val,
                "name": ws.cell(row, 2).value,
                "purpose": ws.cell(row, 3).value,
            })

    module_parameters = []
    if "Module Parameters" in wb.sheetnames:
        ws = wb["Module Parameters"]
        for row in range(3, 200):
            module_id = ws.cell(row, 1).value
            if not module_id:
                break
            module_parameters.append({
                "module_id": module_id,
                "parameter": ws.cell(row, 2).value,
                "value": ws.cell(row, 3).value,
                "post_build": ws.cell(row, 4).value,
            })

    os_tasks = []
    if "Schedule Tables" in wb.sheetnames:
        ws = wb["Schedule Tables"]
        for row in range(3, 100):
            task_id = ws.cell(row, 1).value
            if not task_id:
                break
            os_tasks.append({
                "id": task_id,
                "name": ws.cell(row, 2).value,
                "period_ms": ws.cell(row, 3).value,
                "priority": ws.cell(row, 4).value,
            })

    memory_regions = []
    if "Memory Map" in wb.sheetnames:
        ws = wb["Memory Map"]
        for row in range(3, 100):
            region = ws.cell(row, 1).value
            if not region:
                break
            memory_regions.append({
                "region": region,
                "start": ws.cell(row, 2).value,
                "size_kb": ws.cell(row, 3).value,
                "purpose": ws.cell(row, 4).value,
            })

    return ProbedBSW(
        title=title,
        ecu_target=ecu_target,
        mcal_modules=mcal_modules,
        ecu_abstraction=ecu_abstraction,
        service_layer=service_layer,
        complex_drivers=complex_drivers,
        module_parameters=module_parameters,
        os_tasks=os_tasks,
        memory_regions=memory_regions,
    )
