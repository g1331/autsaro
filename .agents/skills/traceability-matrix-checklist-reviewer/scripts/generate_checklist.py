"""Traceability Matrix checklist reviewer."""

from __future__ import annotations

import sys
from datetime import date
from collections import Counter

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from trace_probe import probe_trace_matrix
from check_definitions import CHECKS, evaluate_trace_matrix
from dashboard import build_dashboard

NAVY = "1F3864"
LIGHT_BLUE = "D9E2F3"
GREY = "D0CECE"
GREEN_OK = "C6EFCE"
RED_BAD = "F8CBAD"
ORANGE_PARTIAL = "FFD966"

def main(matrix_xlsx: str, output_xlsx: str):
    try:
        probed = probe_trace_matrix(matrix_xlsx)
        print(f"Probed Traceability Matrix: {len(probed.elements)} elements")
    except Exception as e:
        print(f"Error probing Traceability Matrix: {e}")
        sys.exit(1)

    results = evaluate_trace_matrix(probed)

    wb = Workbook()
    wb.remove(wb.active)

    # Title
    ws = wb.create_sheet("Title")
    ws.merge_cells("A1:D1")
    c = ws["A1"]
    c.value = "Traceability Matrix Confirmation Measures Checklist"
    c.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="center")

    row = 4
    for label, value in [("Item", "Traceability Matrix"), ("Elements", len(probed.elements)),
                         ("Checklist Generated", date.today().isoformat()),
                         ("Standard", "ISO 26262, ISO 21434, IEEE 1012")]:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        row += 1

    # Checklist
    ws = wb.create_sheet("Checklist")
    ws["A1"] = "Traceability Matrix Confirmation Measures"
    ws["A1"].font = Font(size=12, bold=True)
    ws.merge_cells("A1:D1")

    headers = ["ID", "Confirmation Measure", "Category", "Assessment"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=3, column=col, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center")

    for row_idx, (check_id, result) in enumerate(results.items(), 4):
        check_def = CHECKS.get(check_id)
        if not check_def:
            continue

        ws.cell(row=row_idx, column=1, value=check_id)
        ws.cell(row=row_idx, column=2, value=check_def["measure"])
        ws.cell(row=row_idx, column=3, value=check_def["category"])

        rating = result.get("rating", "NA")
        c = ws.cell(row=row_idx, column=4, value=rating)
        color_map = {"FC": GREEN_OK, "LC": LIGHT_BLUE, "PC": ORANGE_PARTIAL, "NO": RED_BAD, "NA": GREY}
        c.fill = PatternFill("solid", fgColor=color_map.get(rating, GREY))
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center")

    # Findings
    ws = wb.create_sheet("Findings")
    ws["A1"] = "Findings and Recommendations"
    ws["A1"].font = Font(size=12, bold=True)
    ws.merge_cells("A1:C1")

    row = 3
    failures = [(cid, r) for cid, r in results.items() if r.get("rating") == "NO"]
    if failures:
        ws.cell(row=row, column=1, value="FINDINGS (Blockers)").font = Font(bold=True, color="C00000")
        row += 1
        for col, h in enumerate(["ID", "Confirmation Measure", "Finding", "Evidence"], 1):
            c = ws.cell(row=row, column=col, value=h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=NAVY)
        row += 1
        for cid, f in failures:
            ws.cell(row=row, column=1, value=cid)
            ws.cell(row=row, column=2, value=CHECKS.get(cid, {}).get("measure", ""))
            ws.cell(row=row, column=3, value=f.get("finding", "No details"))
            ws.cell(row=row, column=4, value=f.get("evidence", ""))
            row += 1
    else:
        ws.cell(row=row, column=1, value="No automatic document failures recorded; draft checks and execution evidence still require review.")
    for col, width in [(1, 10), (2, 42), (3, 70), (4, 34)]:
        ws.column_dimensions[chr(64 + col)].width = width

    # Dashboard
    ws = wb.create_sheet("Dashboard")
    ws["A1"] = "Traceability Matrix Assessment Dashboard"
    ws["A1"].font = Font(size=12, bold=True)

    counts = Counter(r.get("rating", "NA") for r in results.values())
    row = 3
    for rating, label in [("FC", "Fully Compliant"), ("LC", "Largely Compliant"),
                          ("PC", "Partially Compliant"), ("NO", "Non-Conforming")]:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=counts.get(rating, 0))
        row += 1

    # Signoff
    ws = wb.create_sheet("Signoff")
    ws["A1"] = "Checklist Signoff"
    ws["A1"].font = Font(size=12, bold=True)

    ws.cell(row=3, column=1, value="Reviewer Name:")
    ws.cell(row=3, column=2, value="TBD")
    ws.cell(row=4, column=1, value="Date:")
    ws.cell(row=4, column=2, value=date.today().isoformat())

    try:
        execution = wb.create_sheet("Execution Evidence")
        execution.append(["Test ID", "Reported result (not independently verified)"])
        for element in probed.elements:
            if element.get("type") == "Test":
                execution.append([element["id"], element.get("result") or "NOT REPORTED"])
        _repository_notice(wb)
        wb.save(output_xlsx)
        print(f"Checklist generated: {output_xlsx}")
    except Exception as e:
        print(f"Error saving checklist: {e}")
        sys.exit(1)



def _repository_notice(wb):
    import json
    from pathlib import Path
    import sys
    for sheet in wb:
        for row in sheet:
            for cell in row:
                if cell.value == "APPROVED":
                    cell.value = "DOCUMENT CHECKS COMPLETE - review required"
                elif cell.value == "CONDITIONAL APPROVAL":
                    cell.value = "DOCUMENT ISSUES - review required"
                elif cell.value in ("Internal review comments incorporated", "Released for architecture review", "Initial release"):
                    cell.value = "TEMPLATE PLACEHOLDER - not a recorded project event"
    if "03_Coding_Rules_Catalog" in wb.sheetnames:
        rules = wb["03_Coding_Rules_Catalog"]
        col = rules.max_column + 1
        rules.cell(1, col, "Repository validation status")
        for row in range(2, rules.max_row + 1):
            rules.cell(row, col, "UNVERIFIED UPSTREAM TEMPLATE - rule ID/text/severity/CWE require authoritative review")
    ws = wb.create_sheet("Repository Scope")
    ws.append(["Field", "Value"])
    ws.append(["Purpose", "Document generation/review only; NOT a runtime test or certification"])
    ws.append(["Evidence", "Separate CONTENT / STRUCTURE / DRAFT; missing evidence remains unassessed"])
    ws.append(["Version", "Upstream AUTOSAR R22-11 templates; repository R24-11 obligations need verification"])
    ws.append(["Thresholds", "Upstream thresholds are advisory, not repository acceptance gates"])
    ws.append(["Templates", "Fixed example/default content requires review; not project facts"])
    if len(sys.argv) > 1 and Path(sys.argv[1]).suffix.lower() == ".json":
        with open(sys.argv[1], encoding="utf-8") as stream:
            provenance = json.load(stream).get("provenance", {})
        for key, value in provenance.items():
            ws.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 100


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(f"Usage: {sys.argv[0]} <matrix.xlsx> <output_checklist.xlsx>")
        sys.exit(1)
    main(sys.argv[1], sys.argv[2])
