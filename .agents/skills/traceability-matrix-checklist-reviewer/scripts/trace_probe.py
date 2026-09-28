"""Probe a Traceability Matrix xlsx and extract structured data.

Sheet-name and column contract with traceability-matrix-builder:

  Trace Source Catalog : row 1 header, then Category | Element ID | Title | Type
  Coverage Analysis    : a metric block whose column A carries the metric name
                         and column B the percentage, and an orphan register
                         whose column A carries "Orphan Requirement" /
                         "Orphan Test Case" / "Orphan Stakeholder Need" and
                         column B the element ID.

Both blocks are located by LABEL, not by fixed row offset, so inserting a row
in the builder does not silently empty the probe (which is how the previous
fixed `min_row=6, max_row=7` window failed).
"""

from dataclasses import dataclass, field

from openpyxl import load_workbook

METRIC_LABELS = {
    "need coverage": "need_coverage",
    "requirement coverage": "requirement_coverage",
    "test coverage": "test_coverage",
    "overall coverage": "overall_coverage",
}

ORPHAN_LABELS = {
    "orphan requirement": "orphan_reqs",
    "orphan test case": "orphan_tests",
    "orphan stakeholder need": "orphan_needs",
}


@dataclass
class ProbedTraceMatrix:
    title: str = "Traceability Matrix"
    doc_id: str = "TM-001"
    revision: str = "1.0"
    sheetnames: list = field(default_factory=list)
    elements: list = field(default_factory=list)
    orphan_reqs: list = field(default_factory=list)
    orphan_tests: list = field(default_factory=list)
    orphan_needs: list = field(default_factory=list)
    metrics: dict = field(default_factory=dict)


def _as_float(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def probe_trace_matrix(xlsx_path: str) -> ProbedTraceMatrix:
    """Extract metadata from Traceability Matrix xlsx."""
    try:
        wb = load_workbook(xlsx_path, data_only=True)
    except Exception as e:
        raise ValueError(f"Cannot open {xlsx_path}: {e}")

    probed = ProbedTraceMatrix()
    probed.sheetnames = list(wb.sheetnames)

    # Title page metadata
    if "Title" in wb.sheetnames:
        ws = wb["Title"]
        for row in ws.iter_rows(min_row=1, max_row=20, min_col=1, max_col=2, values_only=True):
            if not row or not row[0]:
                continue
            label = str(row[0]).strip().lower()
            if label == "document id" and row[1]:
                probed.doc_id = str(row[1])
            elif label == "revision" and row[1]:
                probed.revision = str(row[1])
            elif label == "item" and row[1]:
                probed.title = str(row[1])

    # Trace Source Catalog: category | id | title | type
    if "Trace Source Catalog" in wb.sheetnames:
        ws = wb["Trace Source Catalog"]
        for row in ws.iter_rows(min_row=2, max_row=2000, min_col=1, max_col=6, values_only=True):
            if row and row[0] and row[1]:
                probed.elements.append(
                    {
                        "category": row[0],
                        "id": row[1],
                        "title": row[2] if len(row) > 2 else "",
                        "type": row[3] if len(row) > 3 else "",
                        "links": row[4] if len(row) > 4 else "",
                        "result": row[5] if len(row) > 5 else "",
                    }
                )

    # Coverage Analysis: metric block + orphan register, both located by label
    if "Coverage Analysis" in wb.sheetnames:
        ws = wb["Coverage Analysis"]
        for row in ws.iter_rows(min_row=1, max_row=2000, min_col=1, max_col=4, values_only=True):
            if not row or not row[0]:
                continue
            label = str(row[0]).strip().lower()

            if label in METRIC_LABELS:
                value = _as_float(row[1] if len(row) > 1 else None)
                if value is not None:
                    probed.metrics[METRIC_LABELS[label]] = value
                continue

            if label in ORPHAN_LABELS:
                bucket = getattr(probed, ORPHAN_LABELS[label])
                bucket.append(
                    {
                        "id": row[1] if len(row) > 1 else "",
                        "title": row[2] if len(row) > 2 else "",
                        "reason": row[3] if len(row) > 3 else "",
                    }
                )

    return probed
