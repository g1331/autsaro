"""Secure Coding Guidelines probe - extracts structure from guidelines xlsx."""
from __future__ import annotations
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from openpyxl import load_workbook

@dataclass
class ProbedSecCoding:
    path: str = ""
    title: str = ""
    revision: str = ""
    status: str = ""
    author: str = ""
    approver: str = ""
    doc_id: str = ""
    sheet_names: list[str] = field(default_factory=list)
    scope_defined: bool = False
    coding_rules: int = 0
    memory_safety_rules: int = 0
    integer_safety_rules: int = 0
    string_io_rules: int = 0
    crypto_rules: int = 0
    auth_rules: int = 0
    logging_rules: int = 0
    cwe_coverage: int = 0
    tool_configs: int = 0
    deviation_process: bool = False
    is_builder_format: bool = False

def _norm(s: Any) -> str:
    if s is None: return ""
    return str(s).strip()

def probe(xlsx_path: str | Path) -> ProbedSecCoding:
    wb = load_workbook(xlsx_path, data_only=True)
    p = ProbedSecCoding(path=str(xlsx_path), sheet_names=list(wb.sheetnames))

    p.is_builder_format = "03_Coding_Rules_Catalog" in wb.sheetnames

    # Title page
    if "00_Title_Page" in wb.sheetnames:
        ws = wb["00_Title_Page"]
        for r in range(1, min(ws.max_row + 1, 30)):
            for c in range(1, min(ws.max_column + 1, 6)):
                label = _norm(ws.cell(r, c).value).lower()
                value = _norm(ws.cell(r, c + 1).value)
                if "title" in label: p.title = value
                elif "revision" in label: p.revision = value
                elif "status" in label: p.status = value
                elif "author" in label: p.author = value
                elif "approver" in label: p.approver = value
                elif "doc" in label: p.doc_id = value

    # Scope
    if "02_Scope" in wb.sheetnames:
        ws = wb["02_Scope"]
        p.scope_defined = ws.max_row > 1

    # Coding Rules Catalog
    if "03_Coding_Rules_Catalog" in wb.sheetnames:
        ws = wb["03_Coding_Rules_Catalog"]
        p.coding_rules = sum(1 for r in range(2, ws.max_row + 1) if _norm(ws.cell(r, 1).value))

    # Memory Safety
    if "04_Memory_Safety_Rules" in wb.sheetnames:
        ws = wb["04_Memory_Safety_Rules"]
        p.memory_safety_rules = sum(1 for r in range(2, ws.max_row + 1) if _norm(ws.cell(r, 1).value))

    # Integer Safety
    if "05_Integer_Safety_Rules" in wb.sheetnames:
        ws = wb["05_Integer_Safety_Rules"]
        p.integer_safety_rules = sum(1 for r in range(2, ws.max_row + 1) if _norm(ws.cell(r, 1).value))

    # String/IO
    if "06_String_IO_Rules" in wb.sheetnames:
        ws = wb["06_String_IO_Rules"]
        p.string_io_rules = sum(1 for r in range(2, ws.max_row + 1) if _norm(ws.cell(r, 1).value))

    # Cryptography
    if "07_Cryptography_Rules" in wb.sheetnames:
        ws = wb["07_Cryptography_Rules"]
        p.crypto_rules = sum(1 for r in range(2, ws.max_row + 1) if _norm(ws.cell(r, 1).value))

    # Authentication/Authorization
    if "08_Auth_Authorization_Rules" in wb.sheetnames:
        ws = wb["08_Auth_Authorization_Rules"]
        p.auth_rules = sum(1 for r in range(2, ws.max_row + 1) if _norm(ws.cell(r, 1).value))

    # Logging/Audit
    if "09_Logging_Audit_Rules" in wb.sheetnames:
        ws = wb["09_Logging_Audit_Rules"]
        p.logging_rules = sum(1 for r in range(2, ws.max_row + 1) if _norm(ws.cell(r, 1).value))

    # CWE Top 25
    if "10_CWE_Top25_Coverage" in wb.sheetnames:
        ws = wb["10_CWE_Top25_Coverage"]
        p.cwe_coverage = sum(1 for r in range(2, ws.max_row + 1) if _norm(ws.cell(r, 1).value))

    # Tool Configuration
    if "11_Tool_Configuration" in wb.sheetnames:
        ws = wb["11_Tool_Configuration"]
        p.tool_configs = sum(1 for r in range(2, ws.max_row + 1) if _norm(ws.cell(r, 1).value))

    # Deviation Process
    if "12_Deviation_Process" in wb.sheetnames:
        ws = wb["12_Deviation_Process"]
        p.deviation_process = ws.max_row > 1

    return p

if __name__ == "__main__":
    import sys, json
    p = probe(sys.argv[1])
    print(json.dumps({
        "path": p.path, "title": p.title, "revision": p.revision, "status": p.status,
        "author": p.author, "approver": p.approver, "doc_id": p.doc_id,
        "sheets": p.sheet_names, "coding_rules": p.coding_rules,
        "memory_safety": p.memory_safety_rules, "integer_safety": p.integer_safety_rules,
        "string_io": p.string_io_rules, "crypto": p.crypto_rules,
        "auth": p.auth_rules, "logging": p.logging_rules,
        "cwe_coverage": p.cwe_coverage, "tool_configs": p.tool_configs,
        "is_builder_format": p.is_builder_format
    }, indent=2))
