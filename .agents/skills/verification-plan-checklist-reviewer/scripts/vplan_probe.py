"""Probe a Verification Plan xlsx and extract structured data."""

from __future__ import annotations

from pathlib import Path
from dataclasses import dataclass
from openpyxl import load_workbook

@dataclass
class ProbedVPlan:
    title: str = "Verification Plan"
    doc_id: str = "VP-001"
    revision: str = "1.0"
    project: str = ""
    scope: str = ""
    methods: list = None
    environments: dict = None
    test_levels: list = None
    acceptance_criteria: list = None

    def __post_init__(self):
        if self.methods is None:
            self.methods = []
        if self.environments is None:
            self.environments = {}
        if self.test_levels is None:
            self.test_levels = []
        if self.acceptance_criteria is None:
            self.acceptance_criteria = []

def probe_vplan(xlsx_path: str) -> ProbedVPlan:
    """Extract metadata from Verification Plan xlsx."""
    try:
        wb = load_workbook(xlsx_path, data_only=True)
    except Exception as e:
        raise ValueError(f"Cannot open {xlsx_path}: {e}")

    probed = ProbedVPlan()

    # Extract from Title sheet
    if "Title" in wb.sheetnames:
        ws = wb["Title"]
        for row in ws.iter_rows(min_row=5, max_row=20, min_col=1, max_col=2, values_only=False):
            if row[0] and row[0].value:
                key = str(row[0].value).strip().lower()
                val = row[1].value if row[1] else ""
                if "project" in key:
                    probed.project = val or ""
                elif "doc" in key:
                    probed.doc_id = val or ""
                elif "revision" in key:
                    probed.revision = val or ""

    # Extract scope from Verification Scope sheet
    if "Verification Scope" in wb.sheetnames:
        ws = wb["Verification Scope"]
        for row in ws.iter_rows(min_row=6, max_row=10, values_only=True):
            if row and row[0]:
                probed.scope = str(ws["A7"].value or "")
                break

    # Extract methods from Verification Strategy
    if "Verification Strategy" in wb.sheetnames:
        ws = wb["Verification Strategy"]
        for row in ws.iter_rows(min_row=4, max_row=20, min_col=1, max_col=3, values_only=True):
            if row and row[0]:
                probed.methods.append({
                    "method": row[0],
                    "applicability": row[1] if len(row) > 1 else "",
                    "phase": row[2] if len(row) > 2 else ""
                })

    # Extract test levels
    if "Test Levels" in wb.sheetnames:
        ws = wb["Test Levels"]
        for row in ws.iter_rows(min_row=4, max_row=20, min_col=1, max_col=3, values_only=True):
            if row and row[0]:
                probed.test_levels.append({
                    "level": row[0],
                    "description": row[1] if len(row) > 1 else "",
                    "scope": row[2] if len(row) > 2 else ""
                })

    # Extract acceptance criteria
    if "Acceptance Criteria" in wb.sheetnames:
        ws = wb["Acceptance Criteria"]
        for row in ws.iter_rows(min_row=4, max_row=20, min_col=1, max_col=3, values_only=True):
            if row and row[0]:
                probed.acceptance_criteria.append({
                    "criterion": row[0],
                    "measurement": row[1] if len(row) > 1 else "",
                    "threshold": row[2] if len(row) > 2 else ""
                })

    return probed
