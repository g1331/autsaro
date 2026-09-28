"""DTC Catalog probe — reads xlsx and extracts DTC content."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook


@dataclass
class ProbedDTC:
    path: str = ""
    sheet_names: list[str] = field(default_factory=list)
    title: str = ""
    project: str = ""
    doc_id: str = ""
    revision: str = ""
    author: str = ""
    approver: str = ""
    status: str = ""
    has_document_control: bool = False
    revision_history_rows: int = 0
    dtc_count: int = 0
    dtc_codes: list[str] = field(default_factory=list)
    status_mask_bits: int = 0
    debounce_algorithms: list[str] = field(default_factory=list)
    healing_logic_coverage: float = 0.0
    snapshot_bindings_count: int = 0
    extended_data_count: int = 0
    obd_mappings: int = 0
    has_table_of_contents: bool = False


def probe(xlsx_path: str) -> ProbedDTC:
    """Probe a DTC catalog xlsx and return structured data."""
    p = ProbedDTC(path=xlsx_path)

    try:
        wb = load_workbook(xlsx_path, data_only=True)
        p.sheet_names = wb.sheetnames

        # Title page
        if "Title" in wb.sheetnames:
            ws = wb["Title"]
            p.title = ws.cell(3, 1).value or ""
            for row in ws.iter_rows(min_row=6, max_row=15):
                label = row[1].value
                value = row[2].value
                if label == "Project":
                    p.project = value or ""
                elif label == "Document ID":
                    p.doc_id = value or ""
                elif label == "Revision":
                    p.revision = value or ""
                elif label == "Author":
                    p.author = value or ""
                elif label == "Approver":
                    p.approver = value or ""

        # Document Control
        if "Document Control" in wb.sheetnames:
            p.has_document_control = True
            ws = wb["Document Control"]
            for row in ws.iter_rows(min_row=5):
                if row[0].value:
                    p.revision_history_rows += 1

        # DTC Inventory
        if "DTC Inventory" in wb.sheetnames:
            ws = wb["DTC Inventory"]
            for row in ws.iter_rows(min_row=5):
                code = row[0].value
                if code and str(code).strip():
                    p.dtc_codes.append(str(code))
            p.dtc_count = len(p.dtc_codes)

        # Status Mask
        if "Status Mask Configuration" in wb.sheetnames:
            ws = wb["Status Mask Configuration"]
            for row in ws.iter_rows(min_row=5):
                if row[0].value:
                    p.status_mask_bits += 1

        # Debounce
        if "Debounce Algorithms" in wb.sheetnames:
            ws = wb["Debounce Algorithms"]
            for row in ws.iter_rows(min_row=5):
                alg_id = row[0].value
                if alg_id:
                    p.debounce_algorithms.append(str(alg_id))

        # Healing Logic
        if "Healing Logic" in wb.sheetnames:
            ws = wb["Healing Logic"]
            healing_count = 0
            for row in ws.iter_rows(min_row=5):
                if row[0].value and row[1].value:
                    healing_count += 1
            p.healing_logic_coverage = healing_count / max(1, p.dtc_count)

        # Snapshot DIDs
        if "Snapshot DID Bindings" in wb.sheetnames:
            ws = wb["Snapshot DID Bindings"]
            for row in ws.iter_rows(min_row=5):
                if row[0].value:
                    p.snapshot_bindings_count += 1

        # Extended Data
        if "Extended Data Records" in wb.sheetnames:
            ws = wb["Extended Data Records"]
            for row in ws.iter_rows(min_row=5):
                if row[0].value:
                    p.extended_data_count += 1

        # OBD Mapping
        if "OBD Mapping" in wb.sheetnames:
            ws = wb["OBD Mapping"]
            for row in ws.iter_rows(min_row=5):
                if row[0].value:
                    p.obd_mappings += 1

        # TOC check
        p.has_table_of_contents = "Title" in wb.sheetnames or len(p.sheet_names) >= 5

    except Exception:
        pass

    return p
