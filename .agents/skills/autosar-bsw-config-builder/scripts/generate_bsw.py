"""
AUTOSAR BSW Config Builder — generates a multi-tab AUTOSAR Classic BSW configuration workbook.

Usage:
    python generate_bsw.py <input.json> <output.xlsx>

Input JSON schema:
{
  "item": {name, abbr, project, doc_id, revision, date, author, approver, company},
  "ecu_target": {microcontroller, vendor, memory_ram_kb, memory_rom_kb},
  "bus_interfaces": [{name, type, baudrate, channels}],  // CAN, LIN, Ethernet, FlexRay
  "mcal_modules": [{id, name, configured_for_target}],
  "ecu_abstraction": [{id, name, depends_on}],
  "service_layer": [{id, name, critical}],
  "complex_drivers": [{id, name, purpose}],
  "module_parameters": [{module_id, parameter, value, post_build_variant}],
  "memory_layout": [{region, start_address, size_kb, purpose}],
  "os_tasks": [{id, name, period_ms, priority}]
}
"""

from __future__ import annotations

import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT_NAME = "Calibri"
NAVY = "1F3864"
LIGHT_BLUE = "D9E2F3"
ALT_ROW = "F2F2F2"
THIN = Side(border_style="thin", color="BFBFBF")
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def style_title(ws, row, last_col, text):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(name=FONT_NAME, size=14, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 28

def style_header(ws, row, headers):
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_ALL
    ws.row_dimensions[row].height = 36

def build_title(wb, item):
    ws = wb.create_sheet("Title", 0)
    ws.sheet_view.showGridLines = False
    style_title(ws, 2, 4, "AUTOSAR Classic BSW Configuration Workbook")
    style_title(ws, 3, 4, item.get("name", "Unnamed ECU"))
    row = 5
    fields = [
        ("Project", item.get("project")),
        ("Document ID", item.get("doc_id")),
        ("Revision", item.get("revision", "1.0")),
        ("Date", item.get("date", str(date.today()))),
        ("Author", item.get("author")),
        ("Approver", item.get("approver")),
        ("Company", item.get("company")),
    ]
    for label, value in fields:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        row += 1

def build_document_control(wb):
    ws = wb.create_sheet("Document Control", 1)
    style_title(ws, 1, 5, "Document Control")
    headers = ["Change No.", "Revision", "Date", "Author", "Description"]
    style_header(ws, 2, headers)
    ws.cell(row=3, column=1, value=1)
    ws.cell(row=3, column=2, value="1.0")
    ws.cell(row=3, column=3, value=str(date.today()))

def build_ecu_toplevel(wb, ecu_target):
    ws = wb.create_sheet("ECU-C Top-Level", 2)
    style_title(ws, 1, 6, "ECU-C Top-Level Configuration")
    row = 3
    fields = [
        ("Microcontroller", ecu_target.get("microcontroller")),
        ("Vendor", ecu_target.get("vendor")),
        ("RAM (KB)", ecu_target.get("memory_ram_kb")),
        ("ROM (KB)", ecu_target.get("memory_rom_kb")),
    ]
    for label, value in fields:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        row += 1

def build_mcal_inventory(wb, mcal_modules):
    ws = wb.create_sheet("MCAL Module Inventory", 4)
    style_title(ws, 1, 4, "MCAL Module Inventory")
    headers = ["Module ID", "Module Name", "Target Configured", "Port Count"]
    style_header(ws, 2, headers)
    for idx, module in enumerate(mcal_modules, start=3):
        ws.cell(row=idx, column=1, value=module.get("id"))
        ws.cell(row=idx, column=2, value=module.get("name"))
        ws.cell(row=idx, column=3, value=module.get("configured_for_target", "Yes"))
        ws.cell(row=idx, column=4, value="")

def build_ecu_abstraction(wb, ecu_abs):
    ws = wb.create_sheet("ECU Abstraction Inventory", 5)
    style_title(ws, 1, 4, "ECU Abstraction Inventory")
    headers = ["ID", "Module Name", "Depends On", "Complexity"]
    style_header(ws, 2, headers)
    for idx, module in enumerate(ecu_abs, start=3):
        ws.cell(row=idx, column=1, value=module.get("id"))
        ws.cell(row=idx, column=2, value=module.get("name"))
        ws.cell(row=idx, column=3, value=module.get("depends_on", ""))

def build_service_layer(wb, services):
    ws = wb.create_sheet("Service Layer Modules", 6)
    style_title(ws, 1, 4, "Service Layer Modules")
    headers = ["ID", "Module Name", "Critical", "Features"]
    style_header(ws, 2, headers)
    for idx, svc in enumerate(services, start=3):
        ws.cell(row=idx, column=1, value=svc.get("id"))
        ws.cell(row=idx, column=2, value=svc.get("name"))
        ws.cell(row=idx, column=3, value="Yes" if svc.get("critical") else "No")

def build_complex_drivers(wb, drivers):
    ws = wb.create_sheet("Complex Drivers", 7)
    style_title(ws, 1, 4, "Complex Driver Modules")
    headers = ["ID", "Name", "Purpose", "BSW Interaction"]
    style_header(ws, 2, headers)
    for idx, drv in enumerate(drivers, start=3):
        ws.cell(row=idx, column=1, value=drv.get("id"))
        ws.cell(row=idx, column=2, value=drv.get("name"))
        ws.cell(row=idx, column=3, value=drv.get("purpose", ""))

def build_bus_interfaces(wb, bus_interfaces):
    ws = wb.create_sheet("Bus Interfaces", 3)
    style_title(ws, 1, 5, "Bus Interfaces")
    headers = ["Interface", "Type", "Baudrate (bit/s)", "Channels", "Notes"]
    style_header(ws, 2, headers)
    for idx, bus in enumerate(bus_interfaces, start=3):
        ws.cell(row=idx, column=1, value=bus.get("name"))
        ws.cell(row=idx, column=2, value=bus.get("type"))
        ws.cell(row=idx, column=3, value=bus.get("baudrate"))
        ws.cell(row=idx, column=4, value=bus.get("channels"))
        ws.cell(row=idx, column=5, value=bus.get("notes", ""))

def build_module_params(wb, params):
    ws = wb.create_sheet("Module Parameters", 8)
    style_title(ws, 1, 6, "Module Configuration Parameters")
    headers = ["Module ID", "Parameter", "Value", "Post-Build Variant", "Unit", "Notes"]
    style_header(ws, 2, headers)
    for idx, param in enumerate(params, start=3):
        ws.cell(row=idx, column=1, value=param.get("module_id"))
        ws.cell(row=idx, column=2, value=param.get("parameter"))
        ws.cell(row=idx, column=3, value=param.get("value"))
        ws.cell(row=idx, column=4, value=param.get("post_build_variant", ""))

def build_post_build_variants(wb, params, module_names=None):
    """Aggregate the distinct post_build_variant values carried by module_parameters.

    The variant data was always present in `module_parameters[].post_build_variant`;
    before this it was written per-parameter and never rolled up, leaving this tab
    header-only (issue #54). Module IDs are resolved to module names where the
    inventories supply a mapping, so "BSW Modules Affected" reads as BSW module
    names rather than opaque row identifiers.
    """
    ws = wb.create_sheet("Post-Build Variants", 9)
    style_title(ws, 1, 5, "Post-Build Variants")
    headers = ["Variant ID", "Name", "BSW Modules Affected", "Config Set", "Selectable"]
    style_header(ws, 2, headers)

    module_names = module_names or {}
    variants: dict[str, list[str]] = {}
    for param in params:
        variant = param.get("post_build_variant")
        if not variant:
            continue
        module_id = param.get("module_id")
        module = module_names.get(module_id, module_id)
        affected = variants.setdefault(str(variant), [])
        if module and module not in affected:
            affected.append(str(module))

    for idx, (variant, affected) in enumerate(sorted(variants.items()), start=3):
        ws.cell(row=idx, column=1, value=f"PBV{idx - 2:03d}")
        ws.cell(row=idx, column=2, value=variant)
        ws.cell(row=idx, column=3, value=", ".join(affected))
        ws.cell(row=idx, column=4, value=variant)
        ws.cell(row=idx, column=5, value="Yes")

def build_memory_map(wb, memory_layout):
    """Write the `memory_layout` input to the Memory Map tab.

    This function previously took no data argument, so four documented input
    fields were accepted and silently discarded (issue #54). The drop also
    laundered into a passing audit: the reviewer's C005 (Mandatory) rated NA
    "No memory regions defined" against input that did define them.
    """
    ws = wb.create_sheet("Memory Map", 10)
    style_title(ws, 1, 6, "Memory Layout (RAM/ROM Allocation)")
    headers = ["Region", "Start Address", "Size (KB)", "Purpose", "Module Owner", "Access"]
    style_header(ws, 2, headers)
    for idx, region in enumerate(memory_layout, start=3):
        ws.cell(row=idx, column=1, value=region.get("region"))
        ws.cell(row=idx, column=2, value=region.get("start_address"))
        ws.cell(row=idx, column=3, value=region.get("size_kb"))
        ws.cell(row=idx, column=4, value=region.get("purpose"))
        ws.cell(row=idx, column=5, value=region.get("module_owner", ""))
        ws.cell(row=idx, column=6, value=region.get("access", ""))

def build_schedule_tables(wb, os_tasks):
    ws = wb.create_sheet("Schedule Tables", 11)
    style_title(ws, 1, 6, "OS Task Scheduling")
    headers = ["Task ID", "Name", "Period (ms)", "Priority", "State", "Events"]
    style_header(ws, 2, headers)
    for idx, task in enumerate(os_tasks, start=3):
        ws.cell(row=idx, column=1, value=task.get("id"))
        ws.cell(row=idx, column=2, value=task.get("name"))
        ws.cell(row=idx, column=3, value=task.get("period_ms"))
        ws.cell(row=idx, column=4, value=task.get("priority"))

def build_dependencies(wb):
    ws = wb.create_sheet("Inter-Module Dependencies", 12)
    style_title(ws, 1, 5, "Inter-Module Dependencies")
    headers = ["Module A", "Module B", "Dependency Type", "Criticality", "Notes"]
    style_header(ws, 2, headers)

def build_validation_rules(wb):
    ws = wb.create_sheet("Validation Rules", 13)
    style_title(ws, 1, 5, "Validation Rules")
    headers = ["Rule ID", "Rule Description", "Applies To", "Severity", "Status"]
    style_header(ws, 2, headers)

def build_references(wb):
    ws = wb.create_sheet("References", 14)
    style_title(ws, 1, 3, "References and Standards")
    row = 3
    refs = [
        ("AUTOSAR R22-11", "AUTOSAR specification release 22.11"),
        ("BSW Methodology", "See references/methodology.md in this skill"),
        ("Target ECU Documentation", "Microcontroller datasheet"),
    ]
    for ref, desc in refs:
        ws.cell(row=row, column=1, value=ref).font = Font(bold=True)
        ws.cell(row=row, column=2, value=desc)
        row += 1

def main():
    if len(sys.argv) < 2:
        input_json = {
            "item": {"name": "Example ECU", "project": "Project X", "doc_id": "DOC-001", "revision": "1.0"},
            "ecu_target": {"microcontroller": "STM32H743", "vendor": "STMicroelectronics", "memory_ram_kb": 512, "memory_rom_kb": 2048},
            "bus_interfaces": [],
            "mcal_modules": [{"id": "M001", "name": "Port", "configured_for_target": "Yes"}],
            "ecu_abstraction": [{"id": "EA001", "name": "EcuM", "depends_on": ""}],
            "service_layer": [{"id": "SL001", "name": "Com", "critical": True}],
            "complex_drivers": [],
            "module_parameters": [],
            "memory_layout": [],
            "os_tasks": [{"id": "T001", "name": "Task_10ms", "period_ms": 10, "priority": 1}]
        }
    else:
        with open(sys.argv[1], encoding="utf-8") as f:
            input_json = json.load(f)

    wb = Workbook()
    wb.remove(wb.active)

    # module_id -> module name, across every inventory that declares one. Used by
    # the Post-Build Variants roll-up so it names modules rather than row IDs.
    module_names = {
        module.get("id"): module.get("name")
        for key in ("mcal_modules", "ecu_abstraction", "service_layer", "complex_drivers")
        for module in input_json.get(key, [])
        if module.get("id")
    }

    build_title(wb, input_json.get("item", {}))
    build_document_control(wb)
    build_ecu_toplevel(wb, input_json.get("ecu_target", {}))
    build_bus_interfaces(wb, input_json.get("bus_interfaces", []))
    build_mcal_inventory(wb, input_json.get("mcal_modules", []))
    build_ecu_abstraction(wb, input_json.get("ecu_abstraction", []))
    build_service_layer(wb, input_json.get("service_layer", []))
    build_complex_drivers(wb, input_json.get("complex_drivers", []))
    build_module_params(wb, input_json.get("module_parameters", []))
    build_post_build_variants(wb, input_json.get("module_parameters", []), module_names)
    build_memory_map(wb, input_json.get("memory_layout", []))
    build_schedule_tables(wb, input_json.get("os_tasks", []))
    build_dependencies(wb)
    build_validation_rules(wb)
    build_references(wb)

    output_file = sys.argv[2] if len(sys.argv) > 2 else "bsw_config.xlsx"
    _repository_notice(wb)
    wb.save(output_file)
    print(f"Generated {output_file}")



def _repository_notice(wb):
    import json
    from pathlib import Path
    import sys
    for sheet in wb:
        for row in sheet:
            for cell in row:
                if cell.value == "APPROVED":
                    cell.value = "DOCUMENT CHECKS COMPLETE - review required"
                elif cell.value == "CONDITIONAL APPROVAL":
                    cell.value = "DOCUMENT ISSUES - review required"
                elif cell.value in ("Internal review comments incorporated", "Released for architecture review", "Initial release"):
                    cell.value = "TEMPLATE PLACEHOLDER - not a recorded project event"
    if "03_Coding_Rules_Catalog" in wb.sheetnames:
        rules = wb["03_Coding_Rules_Catalog"]
        col = rules.max_column + 1
        rules.cell(1, col, "Repository validation status")
        for row in range(2, rules.max_row + 1):
            rules.cell(row, col, "UNVERIFIED UPSTREAM TEMPLATE - rule ID/text/severity/CWE require authoritative review")
    ws = wb.create_sheet("Repository Scope")
    ws.append(["Field", "Value"])
    ws.append(["Purpose", "Document generation/review only; NOT a runtime test or certification"])
    ws.append(["Evidence", "Separate CONTENT / STRUCTURE / DRAFT; missing evidence remains unassessed"])
    ws.append(["Version", "Upstream AUTOSAR R22-11 templates; repository R24-11 obligations need verification"])
    ws.append(["Thresholds", "Upstream thresholds are advisory, not repository acceptance gates"])
    ws.append(["Templates", "Fixed example/default content requires review; not project facts"])
    if len(sys.argv) > 1 and Path(sys.argv[1]).suffix.lower() == ".json":
        with open(sys.argv[1], encoding="utf-8") as stream:
            provenance = json.load(stream).get("provenance", {})
        for key, value in provenance.items():
            ws.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 100


if __name__ == "__main__":
    main()
