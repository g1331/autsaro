"""AUTOSAR Composition Checklist Reviewer — orchestrator + xlsx writer."""

from __future__ import annotations

import json
import sys
from collections import Counter
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from composition_probe import probe
from check_definitions import CHECKS, CheckResult, execute_checks
from dashboard import build_dashboard


FONT_NAME = "Calibri"
NAVY = "1F3864"
LIGHT_BLUE = "D9E2F3"
ALT_ROW = "F2F2F2"
WARN_YELLOW = "FFF2CC"
GREEN_OK = "C6EFCE"
RED_BAD = "F8CBAD"
ORANGE_PARTIAL = "FFD966"
GREY = "D0CECE"
THIN = Side(border_style="thin", color="BFBFBF")
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_title(ws, row, last_col, text):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(name=FONT_NAME, size=14, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 28


def style_header(ws, row, headers):
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_ALL
    ws.row_dimensions[row].height = 36


def color_assessment(cell, rating, obligation):
    color_map = {"Pass": GREEN_OK, "Partial": ORANGE_PARTIAL, "Fail": RED_BAD, "NA": GREY}
    color = color_map.get(rating)
    if color:
        cell.fill = PatternFill("solid", fgColor=color)
    cell.font = Font(name=FONT_NAME, bold=True, size=10, color="C00000" if rating == "Fail" else "000000")
    cell.alignment = Alignment(horizontal="center", vertical="center")


def build_title(wb, p):
    ws = wb.create_sheet("Title")
    ws.sheet_view.showGridLines = False
    style_title(ws, 2, 5, "AUTOSAR Composition Specification Checklist Review")
    style_title(ws, 3, 5, p.composition_name or "(Composition not detected)")
    fields = [
        ("Project", p.project or ""),
        ("Composition Name", p.composition_name or ""),
        ("Document ID", p.doc_id or ""),
        ("Revision", p.revision or ""),
        ("Source File", p.path or ""),
        ("Checklist Generated", date.today().isoformat()),
        ("Standard", "AUTOSAR R22-11 (Classic Platform)"),
        ("Reviewer", "(to be assigned)"),
    ]
    row = 6
    for label, value in fields:
        ws.cell(row=row, column=2, value=label).font = Font(name=FONT_NAME, size=11, bold=True, color=NAVY)
        ws.cell(row=row, column=3, value=value).font = Font(name=FONT_NAME, size=10)
        ws.cell(row=row, column=3).alignment = Alignment(horizontal="left", wrap_text=True)
        row += 1
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 80


def build_summary(wb, results):
    ws = wb.create_sheet("Summary")
    ws.sheet_view.showGridLines = False
    style_title(ws, 1, 3, "Review Summary")

    status_counts = Counter(r.status for r in results)
    ws.cell(row=3, column=1, value="Status").font = Font(name=FONT_NAME, bold=True, size=11, color=NAVY)
    ws.cell(row=3, column=2, value="Count")

    row = 4
    for status in ["Pass", "Partial", "Fail", "NA"]:
        count = status_counts.get(status, 0)
        ws.cell(row=row, column=1, value=status)
        cell = ws.cell(row=row, column=2, value=count)
        color_assessment(cell, status, "MUST")
        row += 1

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 12


def build_findings(wb, results):
    ws = wb.create_sheet("Findings")
    ws.sheet_view.showGridLines = False
    style_title(ws, 1, 5, "Detailed Check Findings")

    style_header(ws, 3, ["Check ID", "Title", "Status", "Finding", "Remediation"])

    row = 4
    for r in results:
        ws.cell(row=row, column=1, value=r.check_id)
        ws.cell(row=row, column=2, value=r.title)
        cell = ws.cell(row=row, column=3, value=r.status)
        color_assessment(cell, r.status, "MUST")
        ws.cell(row=row, column=4, value=r.finding)
        ws.cell(row=row, column=5, value=r.remediation)
        row += 1

    for col in range(1, 6):
        for r in range(4, row):
            ws.cell(row=r, column=col).border = BORDER_ALL
            ws.cell(row=r, column=col).font = Font(name=FONT_NAME, size=9)
            ws.cell(row=r, column=col).alignment = Alignment(vertical="top", wrap_text=True)

    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 40
    ws.column_dimensions["E"].width = 40


def generate(input_path: str, output_path: str) -> dict:
    """Main checklist generation."""
    p = probe(input_path)
    results = execute_checks(p)

    wb = Workbook()
    wb.remove(wb.active)

    build_title(wb, p)
    build_summary(wb, results)
    build_findings(wb, results)
    build_dashboard(wb, results)

    wb.active = 0
    _repository_notice(wb)
    wb.save(output_path)

    status_counts = Counter(r.status for r in results)
    return {
        "composition_name": p.composition_name or "Unknown",
        "doc_id": p.doc_id or "",
        "checks_total": len(results),
        "checks_pass": status_counts.get("Pass", 0),
        "checks_partial": status_counts.get("Partial", 0),
        "checks_fail": status_counts.get("Fail", 0),
        "checks_na": status_counts.get("NA", 0),
        "output_path": output_path,
    }




def _repository_notice(wb):
    import json
    from pathlib import Path
    import sys
    for sheet in wb:
        for row in sheet:
            for cell in row:
                if cell.value == "APPROVED":
                    cell.value = "DOCUMENT CHECKS COMPLETE - review required"
                elif cell.value == "CONDITIONAL APPROVAL":
                    cell.value = "DOCUMENT ISSUES - review required"
                elif cell.value in ("Internal review comments incorporated", "Released for architecture review", "Initial release"):
                    cell.value = "TEMPLATE PLACEHOLDER - not a recorded project event"
    if "03_Coding_Rules_Catalog" in wb.sheetnames:
        rules = wb["03_Coding_Rules_Catalog"]
        col = rules.max_column + 1
        rules.cell(1, col, "Repository validation status")
        for row in range(2, rules.max_row + 1):
            rules.cell(row, col, "UNVERIFIED UPSTREAM TEMPLATE - rule ID/text/severity/CWE require authoritative review")
    ws = wb.create_sheet("Repository Scope")
    ws.append(["Field", "Value"])
    ws.append(["Purpose", "Document generation/review only; NOT a runtime test or certification"])
    ws.append(["Evidence", "Separate CONTENT / STRUCTURE / DRAFT; missing evidence remains unassessed"])
    ws.append(["Version", "Upstream AUTOSAR R22-11 templates; repository R24-11 obligations need verification"])
    ws.append(["Thresholds", "Upstream thresholds are advisory, not repository acceptance gates"])
    ws.append(["Templates", "Fixed example/default content requires review; not project facts"])
    if len(sys.argv) > 1 and Path(sys.argv[1]).suffix.lower() == ".json":
        with open(sys.argv[1], encoding="utf-8") as stream:
            provenance = json.load(stream).get("provenance", {})
        for key, value in provenance.items():
            ws.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 100


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python generate_checklist.py <input.xlsx> <output.xlsx>", file=sys.stderr)
        sys.exit(2)
    summary = generate(sys.argv[1], sys.argv[2])
    print(json.dumps(summary, indent=2))
