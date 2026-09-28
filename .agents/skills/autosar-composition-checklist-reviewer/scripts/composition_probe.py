"""AUTOSAR Composition workbook probe."""

from openpyxl import load_workbook


class ProbedComposition:
    def __init__(self):
        self.path = None
        self.title = None
        self.project = None
        self.doc_id = None
        self.revision = None
        self.composition_name = None
        self.swc_instances = []
        self.assembly_connectors = []
        self.delegation_connectors = []
        self.boundary_ports = []


def probe(file_path: str) -> ProbedComposition:
    """Extract composition metadata and structure."""
    p = ProbedComposition()
    p.path = file_path

    wb = load_workbook(file_path, data_only=True)

    # Title page
    if "00_Title_Page" in wb.sheetnames:
        ws = wb["00_Title_Page"]
        p.title = ws["A3"].value or "(Not detected)"
        for row in ws.iter_rows(min_row=6, max_row=20, min_col=2, max_col=3):
            if row[0].value and row[1].value:
                label = str(row[0].value).strip()
                value = row[1].value
                if "Project" in label:
                    p.project = value
                elif "Composition Name" in label:
                    p.composition_name = value
                elif "Document ID" in label:
                    p.doc_id = value
                elif "Revision" in label:
                    p.revision = value

    # SWC instances
    if "03_SWC_Instances" in wb.sheetnames:
        ws = wb["03_SWC_Instances"]
        for row in ws.iter_rows(min_row=4, max_col=5):
            if row[0].value and row[0].value != "Instance ID":
                p.swc_instances.append({
                    "id": row[0].value,
                    "name": row[1].value,
                    "swc_type": row[2].value,
                    "role": row[3].value,
                    "multiplicity": row[4].value,
                })

    # Assembly connectors
    if "04_Assembly_Connectors" in wb.sheetnames:
        ws = wb["04_Assembly_Connectors"]
        for row in ws.iter_rows(min_row=4, max_col=6):
            if row[0].value and row[0].value != "Connector ID":
                p.assembly_connectors.append({
                    "id": row[0].value,
                    "source_swc": row[1].value,
                    "source_port": row[2].value,
                    "target_swc": row[3].value,
                    "target_port": row[4].value,
                    "interface": row[5].value,
                })

    # Delegation connectors
    if "05_Delegation_Connectors" in wb.sheetnames:
        ws = wb["05_Delegation_Connectors"]
        for row in ws.iter_rows(min_row=4, max_col=5):
            if row[0].value and row[0].value != "Delegate ID":
                p.delegation_connectors.append({
                    "id": row[0].value,
                    "swc_instance": row[1].value,
                    "swc_port": row[2].value,
                    "boundary_port": row[3].value,
                    "direction": row[4].value,
                })

    # Boundary ports
    if "06_Boundary_Ports" in wb.sheetnames:
        ws = wb["06_Boundary_Ports"]
        for row in ws.iter_rows(min_row=4, max_col=5):
            if row[0].value and row[0].value != "Port ID":
                p.boundary_ports.append({
                    "id": row[0].value,
                    "name": row[1].value,
                    "category": row[2].value,
                    "direction": row[3].value,
                    "interface": row[4].value,
                })

    return p
