"""Probe AUTOSAR RTE Mapping xlsx and extract key structures."""

from __future__ import annotations

from dataclasses import dataclass
from openpyxl import load_workbook

@dataclass
class ProbedRTE:
    title: str
    ecu_id: str
    data_mappings: list
    runnable_to_task: list
    schedulable_tasks: list
    exclusive_areas: list
    mode_switches: list

def probe(xlsx_path: str) -> ProbedRTE:
    wb = load_workbook(xlsx_path)

    title = ""
    if "Title" in wb.sheetnames:
        ws = wb["Title"]
        title = ws.cell(3, 1).value or "RTE Mapping"

    ecu_id = ""
    if "ECU Integration Identity" in wb.sheetnames:
        ws = wb["ECU Integration Identity"]
        ecu_id = ws.cell(4, 2).value or ""

    data_mappings = []
    if "Data Mappings" in wb.sheetnames:
        ws = wb["Data Mappings"]
        for row in range(3, 200):
            swc = ws.cell(row, 1).value
            if not swc:
                break
            data_mappings.append({
                "swc": swc,
                "port": ws.cell(row, 2).value,
                "port_type": ws.cell(row, 3).value,
                "com_signal": ws.cell(row, 4).value,
                "signal_type": ws.cell(row, 5).value,
            })

    runnable_to_task = []
    if "Runnable-to-Task" in wb.sheetnames:
        ws = wb["Runnable-to-Task"]
        for row in range(3, 200):
            swc = ws.cell(row, 1).value
            if not swc:
                break
            runnable_to_task.append({
                "swc": swc,
                "runnable": ws.cell(row, 2).value,
                "task": ws.cell(row, 3).value,
                "trigger": ws.cell(row, 4).value,
            })

    schedulable_tasks = []
    if "Schedulable Period" in wb.sheetnames:
        ws = wb["Schedulable Period"]
        for row in range(3, 100):
            task_name = ws.cell(row, 1).value
            if not task_name:
                break
            schedulable_tasks.append({
                "name": task_name,
                "period_ms": ws.cell(row, 2).value,
                "priority": ws.cell(row, 3).value,
            })

    exclusive_areas = []
    if "Exclusive Areas" in wb.sheetnames:
        ws = wb["Exclusive Areas"]
        for row in range(3, 100):
            swc = ws.cell(row, 1).value
            if not swc:
                break
            exclusive_areas.append({
                "swc": swc,
                "area": ws.cell(row, 2).value,
                "os_resource": ws.cell(row, 3).value,
            })

    mode_switches = []
    if "Mode-Switch Mapping" in wb.sheetnames:
        ws = wb["Mode-Switch Mapping"]
        for row in range(3, 100):
            swc = ws.cell(row, 1).value
            if not swc:
                break
            mode_switches.append({
                "swc": swc,
                "mode_port": ws.cell(row, 2).value,
                "mode_group": ws.cell(row, 3).value,
            })

    return ProbedRTE(
        title=title,
        ecu_id=ecu_id,
        data_mappings=data_mappings,
        runnable_to_task=runnable_to_task,
        schedulable_tasks=schedulable_tasks,
        exclusive_areas=exclusive_areas,
        mode_switches=mode_switches,
    )
