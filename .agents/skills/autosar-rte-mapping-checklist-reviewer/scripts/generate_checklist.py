"""RTE Mapping Checklist Reviewer — orchestrator + xlsx writer."""

from __future__ import annotations

import sys
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from rte_probe import ProbedRTE, probe
from check_definitions import CHECKS, CheckResult
from dashboard import build_dashboard

FONT_NAME = "Calibri"
NAVY = "1F3864"
LIGHT_BLUE = "D9E2F3"
GREEN_OK = "C6EFCE"
RED_BAD = "F8CBAD"
ORANGE_PARTIAL = "FFD966"
GREY = "D0CECE"
THIN = Side(border_style="thin", color="BFBFBF")
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def style_title(ws, row, last_col, text):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(name=FONT_NAME, size=14, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 28

def style_header(ws, row, headers):
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_ALL
    ws.row_dimensions[row].height = 36

def color_assessment(cell, rating):
    color_map = {"FC": GREEN_OK, "LC": LIGHT_BLUE, "PC": ORANGE_PARTIAL, "NO": RED_BAD, "NA": GREY}
    color = color_map.get(rating)
    if color:
        cell.fill = PatternFill("solid", fgColor=color)
    cell.font = Font(name=FONT_NAME, bold=True, size=10, color="C00000" if rating == "NO" else "000000")

def build_title(wb, title):
    ws = wb.create_sheet("Title", 0)
    ws.sheet_view.showGridLines = False
    style_title(ws, 2, 5, "RTE Mapping Confirmation Measures Checklist")
    style_title(ws, 3, 5, title or "(RTE title not detected)")

def build_document_control(wb):
    ws = wb.create_sheet("Document Control", 1)
    style_title(ws, 1, 4, "Document Control")
    headers = ["Change No.", "Revision", "Date", "Description"]
    style_header(ws, 2, headers)
    ws.cell(row=3, column=1, value=1)
    ws.cell(row=3, column=2, value="1.0")
    ws.cell(row=3, column=3, value=str(date.today()))

def build_checklist_tab(wb, tab_name, checks, results):
    ws = wb.create_sheet(tab_name, len(wb.sheetnames))
    style_title(ws, 1, 7, f"{tab_name} Compliance")
    headers = ["Check ID", "Requirement", "Finding", "Rating", "Obligation", "Recommended Action", "Notes"]
    style_header(ws, 2, headers)

    row = 3
    for check, result in zip(checks, results):
        ws.cell(row=row, column=1, value=check.id)
        ws.cell(row=row, column=2, value=check.requirement)
        ws.cell(row=row, column=3, value=result.finding or "")
        cell = ws.cell(row=row, column=4, value=result.rating)
        color_assessment(cell, result.rating)
        ws.cell(row=row, column=5, value=check.obligation)
        ws.cell(row=row, column=6, value=result.recommended_action or "")
        row += 1

def build_references(wb):
    ws = wb.create_sheet("References", len(wb.sheetnames))
    style_title(ws, 1, 3, "References")
    row = 3
    refs = [
        ("AUTOSAR R22-11", "AUTOSAR Classic specification"),
        ("RTE Methodology", "scripts/references/rte_checks.md"),
    ]
    for ref, desc in refs:
        ws.cell(row=row, column=1, value=ref).font = Font(bold=True)
        ws.cell(row=row, column=2, value=desc)
        row += 1

def main():
    if len(sys.argv) < 2:
        print("Usage: python generate_checklist.py <rte_mapping.xlsx> [output.xlsx]")
        sys.exit(1)

    rte_file = sys.argv[1]
    output_file = sys.argv[2] if len(sys.argv) > 2 else "rte_checklist.xlsx"

    probed = probe(rte_file)
    results = {}

    for check in CHECKS:
        check_result = check.execute(probed)
        if check.tab not in results:
            results[check.tab] = []
        results[check.tab].append((check, check_result))

    wb = Workbook()
    wb.remove(wb.active)

    build_title(wb, probed.title)
    build_document_control(wb)
    build_dashboard(wb, results)

    for tab_code in sorted(results.keys()):
        checks = [c for c, _ in results[tab_code]]
        check_results = [r for _, r in results[tab_code]]
        build_checklist_tab(wb, f"Tab_{tab_code}", checks, check_results)

    build_references(wb)
    _repository_notice(wb)
    wb.save(output_file)
    print(f"Generated {output_file}")



def _repository_notice(wb):
    import json
    from pathlib import Path
    import sys
    for sheet in wb:
        for row in sheet:
            for cell in row:
                if cell.value == "APPROVED":
                    cell.value = "DOCUMENT CHECKS COMPLETE - review required"
                elif cell.value == "CONDITIONAL APPROVAL":
                    cell.value = "DOCUMENT ISSUES - review required"
                elif cell.value in ("Internal review comments incorporated", "Released for architecture review", "Initial release"):
                    cell.value = "TEMPLATE PLACEHOLDER - not a recorded project event"
    if "03_Coding_Rules_Catalog" in wb.sheetnames:
        rules = wb["03_Coding_Rules_Catalog"]
        col = rules.max_column + 1
        rules.cell(1, col, "Repository validation status")
        for row in range(2, rules.max_row + 1):
            rules.cell(row, col, "UNVERIFIED UPSTREAM TEMPLATE - rule ID/text/severity/CWE require authoritative review")
    ws = wb.create_sheet("Repository Scope")
    ws.append(["Field", "Value"])
    ws.append(["Purpose", "Document generation/review only; NOT a runtime test or certification"])
    ws.append(["Evidence", "Separate CONTENT / STRUCTURE / DRAFT; missing evidence remains unassessed"])
    ws.append(["Version", "Upstream AUTOSAR R22-11 templates; repository R24-11 obligations need verification"])
    ws.append(["Thresholds", "Upstream thresholds are advisory, not repository acceptance gates"])
    ws.append(["Templates", "Fixed example/default content requires review; not project facts"])
    if len(sys.argv) > 1 and Path(sys.argv[1]).suffix.lower() == ".json":
        with open(sys.argv[1], encoding="utf-8") as stream:
            provenance = json.load(stream).get("provenance", {})
        for key, value in provenance.items():
            ws.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 100


if __name__ == "__main__":
    main()
