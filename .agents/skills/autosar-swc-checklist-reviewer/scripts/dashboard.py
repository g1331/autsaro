"""Visual dashboard for SWC checklist review."""

from collections import Counter
from openpyxl.styles import Alignment, Font, PatternFill


FONT_NAME = "Calibri"
NAVY = "1F3864"
GREEN_OK = "C6EFCE"
RED_BAD = "F8CBAD"
ORANGE_PARTIAL = "FFD966"


def build_dashboard(wb, results):
    """Build visual dashboard tab with summary metrics."""
    ws = wb.create_sheet("Dashboard", 0)
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:E1")
    title = ws["A1"]
    title.value = "SWC Checklist Review Dashboard"
    title.font = Font(name=FONT_NAME, size=16, bold=True, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor=NAVY)
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    status_counts = Counter(r.status for r in results)
    total = len(results)
    passed = status_counts.get("Pass", 0)
    partial = status_counts.get("Partial", 0)
    failed = status_counts.get("Fail", 0)
    na = status_counts.get("NA", 0)
    applicable = max(1, total - na)
    compliant_pct = (passed / applicable * 100) if applicable else 0

    metrics = [
        ("Total Checks", str(total)),
        ("Pass", str(passed)),
        ("Partial", str(partial)),
        ("Fail", str(failed)),
        ("N/A", str(na)),
        ("Compliance %", f"{compliant_pct:.0f}%"),
    ]

    row = 3
    for label, value in metrics:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True, size=11)
        cell = ws.cell(row=row, column=2, value=value)
        cell.font = Font(bold=True, size=12)
        if label == "Pass":
            cell.fill = PatternFill("solid", fgColor=GREEN_OK)
        elif label == "Partial":
            cell.fill = PatternFill("solid", fgColor=ORANGE_PARTIAL)
        elif label == "Fail":
            cell.fill = PatternFill("solid", fgColor=RED_BAD)
        row += 1

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 12
