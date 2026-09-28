"""AUTOSAR SWC workbook probe — extract structure and metadata."""

from pathlib import Path
from openpyxl import load_workbook


class ProbedSWC:
    def __init__(self):
        self.path = None
        self.title = None
        self.project = None
        self.doc_id = None
        self.revision = None
        self.status = None
        self.swc_name = None
        self.swc_category = None
        self.ports = []
        self.runnables = []
        self.events = []
        self.exclusive_areas = []
        self.data_types = []
        self.impl_data_types = []


def probe(file_path: str) -> ProbedSWC:
    """Extract SWC metadata and structure from workbook."""
    p = ProbedSWC()
    p.path = file_path

    wb = load_workbook(file_path, data_only=True)

    # Title page
    if "00_Title_Page" in wb.sheetnames:
        ws = wb["00_Title_Page"]
        p.title = ws["A3"].value or "(Not detected)"
        for row in ws.iter_rows(min_row=6, max_row=20, min_col=2, max_col=3):
            if row[0].value and row[1].value:
                label = str(row[0].value).strip()
                value = row[1].value
                if "Project" in label:
                    p.project = value
                elif "SWC Name" in label:
                    p.swc_name = value
                elif "SWC Category" in label:
                    p.swc_category = value
                elif "Document ID" in label:
                    p.doc_id = value
                elif "Revision" in label:
                    p.revision = value

    # Port inventory
    if "03_Port_Inventory" in wb.sheetnames:
        ws = wb["03_Port_Inventory"]
        for row in ws.iter_rows(min_row=4, max_col=6):
            if row[0].value and row[0].value != "Port ID":
                p.ports.append({
                    "id": row[0].value,
                    "name": row[1].value,
                    "category": row[2].value,
                    "direction": row[3].value,
                    "interface_name": row[4].value,
                    "interface_type": row[5].value,
                })

    # Runnable catalog
    if "08_Runnable_Catalog" in wb.sheetnames:
        ws = wb["08_Runnable_Catalog"]
        current_runnable = {}
        for row in ws.iter_rows(min_row=3):
            cell_val = row[0].value
            if not cell_val:
                continue
            if isinstance(cell_val, str) and cell_val.startswith("ID:"):
                if current_runnable:
                    p.runnables.append(current_runnable)
                current_runnable = {"id": cell_val.replace("ID:", "").strip()}
            elif isinstance(cell_val, str) and "Name:" in cell_val:
                current_runnable["name"] = cell_val.replace("Name:", "").strip()
        if current_runnable:
            p.runnables.append(current_runnable)

    # Event catalog
    if "09_Event_Catalog" in wb.sheetnames:
        ws = wb["09_Event_Catalog"]
        current_event = {}
        for row in ws.iter_rows(min_row=3):
            cell_val = row[0].value
            if not cell_val:
                continue
            if isinstance(cell_val, str) and cell_val.startswith("ID:"):
                if current_event:
                    p.events.append(current_event)
                current_event = {"id": cell_val.replace("ID:", "").strip()}
            elif isinstance(cell_val, str) and "Name:" in cell_val:
                current_event["name"] = cell_val.replace("Name:", "").strip()
        if current_event:
            p.events.append(current_event)

    # Exclusive areas
    if "10_Exclusive_Areas" in wb.sheetnames:
        ws = wb["10_Exclusive_Areas"]
        for row in ws.iter_rows(min_row=4, max_col=3):
            if row[0].value and row[0].value != "Exclusive Area ID":
                p.exclusive_areas.append({
                    "id": row[0].value,
                    "name": row[1].value,
                    "runnables": row[2].value,
                })

    # Data types
    if "11_Data_Types" in wb.sheetnames:
        ws = wb["11_Data_Types"]
        for row in ws.iter_rows(min_row=4, max_col=6):
            if row[0].value and row[0].value != "Type ID":
                p.data_types.append({
                    "id": row[0].value,
                    "name": row[1].value,
                    "base_type": row[2].value,
                })

    # Implementation data types
    if "12_Implementation_Data_Types" in wb.sheetnames:
        ws = wb["12_Implementation_Data_Types"]
        for row in ws.iter_rows(min_row=4, max_col=4):
            if row[0].value and row[0].value != "Impl Type ID":
                p.impl_data_types.append({
                    "id": row[0].value,
                    "name": row[1].value,
                    "category": row[2].value,
                })

    return p
