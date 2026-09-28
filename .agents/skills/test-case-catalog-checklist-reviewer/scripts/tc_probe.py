"""Read both repository test inventories and legacy three-column catalogs."""

from dataclasses import dataclass, field

from openpyxl import load_workbook


@dataclass
class TestCase:
    id: str = ""
    title: str = ""
    precondition: str = ""
    inputs: str = ""
    expected_outputs: str = ""
    req_id: str = ""
    pass_criterion: str = ""
    automation: str = ""
    test_level: str = ""
    environment: str = ""
    priority: str = ""


@dataclass
class ProbedTCCatalog:
    title: str = ""
    doc_id: str = ""
    revision: str = ""
    test_cases: list[TestCase] = field(default_factory=list)


def probe_tc_catalog(xlsx_path):
    wb = load_workbook(xlsx_path, data_only=True)
    probed = ProbedTCCatalog()
    try:
        name = next(
            (n for n in ("Test Case Inventory", "Test Cases") if n in wb.sheetnames),
            None,
        )
        if name is None:
            return probed
        rows = wb[name].iter_rows(values_only=True)
        headers = [str(v or "").strip() for v in next(rows, ())]
        labels = {
            "ID": "id",
            "Title": "title",
            "Precondition": "precondition",
            "Inputs": "inputs",
            "Expected Outputs": "expected_outputs",
            "Requirement ID": "req_id",
            "Pass Criterion": "pass_criterion",
            "Automation": "automation",
            "Test Level": "test_level",
            "Environment": "environment",
            "Priority": "priority",
        }
        for row in rows:
            if not any(value is not None for value in row):
                continue
            values = {
                labels[key]: str(value) if value is not None else ""
                for key, value in zip(headers, row)
                if key in labels
            }
            probed.test_cases.append(TestCase(**values))
        return probed
    finally:
        wb.close()
