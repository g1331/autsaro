"""UDS service probe — extract data from xlsx for analysis."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from openpyxl import load_workbook


@dataclass
class ProbedUDS:
    """Probed state from UDS service catalog xlsx."""
    path: str = ""
    title: str = ""
    project: str = ""
    doc_id: str = ""
    revision: str = ""
    status: str = ""
    author: str = ""
    approver: str = ""

    sheet_names: list[str] = field(default_factory=list)
    has_ecu_identity: bool = False

    logical_address: str = ""
    physical_address: str = ""
    functional_address: str = ""
    hw_part_number: str = ""

    session_count: int = 0
    session_types: set[str] = field(default_factory=set)

    service_count: int = 0
    service_ids: set[str] = field(default_factory=set)
    supported_services: list[dict] = field(default_factory=list)

    did_count: int = 0
    did_ids: set[str] = field(default_factory=set)
    did_duplicates: int = 0

    rid_count: int = 0
    rid_ids: set[str] = field(default_factory=set)
    rid_duplicates: int = 0

    security_levels: int = 0
    security_level_ids: set[int] = field(default_factory=set)

    nrc_count: int = 0
    nrc_codes: set[str] = field(default_factory=set)

    p2_ms: str = ""
    p2_star_ms: str = ""
    s3_ms: str = ""

    revision_history_rows: int = 0


def probe(xlsx_path: str) -> ProbedUDS:
    """Load and probe UDS service catalog workbook."""

    p = ProbedUDS(path=xlsx_path)

    try:
        wb = load_workbook(xlsx_path, data_only=True)
    except Exception as e:
        print(f"Error loading {xlsx_path}: {e}")
        return p

    p.sheet_names = wb.sheetnames

    # Title sheet
    if "Title" in wb.sheetnames:
        ws = wb["Title"]
        for row in ws.iter_rows(min_row=1, max_row=50, values_only=True):
            if row[0] and len(row) > 1:
                cell_str = str(row[0]).lower()
                if "name" in cell_str and row[1]:
                    p.title = str(row[1])
                elif "project" in cell_str and row[1]:
                    p.project = str(row[1])
                elif "doc_id" in cell_str and row[1]:
                    p.doc_id = str(row[1])
                elif "revision" in cell_str and row[1]:
                    p.revision = str(row[1])
                elif "status" in cell_str and row[1]:
                    p.status = str(row[1])
                elif "author" in cell_str and row[1]:
                    p.author = str(row[1])
                elif "approver" in cell_str and row[1]:
                    p.approver = str(row[1])

    # Document Control sheet
    if "Document Control" in wb.sheetnames:
        ws = wb["Document Control"]
        p.revision_history_rows = sum(1 for row in ws.iter_rows(min_row=2, values_only=True) if row[0])

    # ECU Diagnostic Identity
    if "ECU Diagnostic Identity" in wb.sheetnames:
        ws = wb["ECU Diagnostic Identity"]
        p.has_ecu_identity = True
        for row in ws.iter_rows(min_row=1, max_row=20, values_only=True):
            if row[0] and len(row) > 1:
                cell_str = str(row[0]).lower()
                if "logical" in cell_str:
                    p.logical_address = str(row[1])
                elif "physical" in cell_str:
                    p.physical_address = str(row[1])
                elif "functional" in cell_str:
                    p.functional_address = str(row[1])
                elif "hw part" in cell_str:
                    p.hw_part_number = str(row[1])

    # Session Definitions
    if "Session Definitions" in wb.sheetnames:
        ws = wb["Session Definitions"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0]:
                p.session_count += 1
                p.session_types.add(str(row[0]))

    # Service Inventory
    if "Service Inventory" in wb.sheetnames:
        ws = wb["Service Inventory"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0]:
                p.service_count += 1
                p.service_ids.add(str(row[0]))
                p.supported_services.append({
                    "id": str(row[0]),
                    "name": str(row[1]) if len(row) > 1 else "",
                    "sub_functions": str(row[2]) if len(row) > 2 else "",
                })

    # DID Catalog
    if "DID Catalog" in wb.sheetnames:
        ws = wb["DID Catalog"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0]:
                did_id = str(row[0]).strip()
                p.did_count += 1
                if did_id in p.did_ids:
                    p.did_duplicates += 1
                p.did_ids.add(did_id)

    # RID Catalog
    if "RID Catalog" in wb.sheetnames:
        ws = wb["RID Catalog"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0]:
                rid_id = str(row[0]).strip()
                p.rid_count += 1
                if rid_id in p.rid_ids:
                    p.rid_duplicates += 1
                p.rid_ids.add(rid_id)

    # Security Access Configuration
    if "Security Access Configuration" in wb.sheetnames:
        ws = wb["Security Access Configuration"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0]:
                p.security_levels += 1
                try:
                    level_int = int(str(row[0]))
                    p.security_level_ids.add(level_int)
                except ValueError:
                    pass

    # Negative Response Codes
    if "Negative Response Codes" in wb.sheetnames:
        ws = wb["Negative Response Codes"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0]:
                p.nrc_count += 1
                p.nrc_codes.add(str(row[0]))

    # Timing Parameters
    if "Timing Parameters" in wb.sheetnames:
        ws = wb["Timing Parameters"]
        for row in ws.iter_rows(min_row=1, max_row=20, values_only=True):
            if row[0] and len(row) > 1:
                cell_str = str(row[0]).lower()
                if "p2" in cell_str and "*" not in cell_str:
                    p.p2_ms = str(row[1])
                elif "p2*" in cell_str:
                    p.p2_star_ms = str(row[1])
                elif "s3" in cell_str:
                    p.s3_ms = str(row[1])

    return p
