"""DEM content probing utility for best-effort introspection of DEM config workbooks."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


@dataclass
class ProbedDEM:
    """Probed DEM content from xlsx."""
    events: list[dict[str, Any]]
    debounce_config: list[dict[str, Any]]
    memory_config: dict[str, Any]
    snapshot_records: list[dict[str, Any]]
    edr_config: list[dict[str, Any]]
    nvm_mapping: list[dict[str, Any]]


def probe(dem_path: str | Path) -> ProbedDEM:
    """Probe a DEM config workbook and extract content."""
    wb = load_workbook(dem_path)

    events = _probe_events(wb)
    debounce_config = _probe_debounce(wb)
    memory_config = _probe_memory(wb)
    snapshot_records = _probe_snapshots(wb)
    edr_config = _probe_edr(wb)
    nvm_mapping = _probe_nvm(wb)

    return ProbedDEM(
        events=events,
        debounce_config=debounce_config,
        memory_config=memory_config,
        snapshot_records=snapshot_records,
        edr_config=edr_config,
        nvm_mapping=nvm_mapping,
    )


def _probe_events(wb) -> list[dict[str, Any]]:
    if "Event Inventory" not in wb.sheetnames:
        return []

    ws = wb["Event Inventory"]
    events = []

    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[0]:
            events.append({
                "event_id": row[0],
                "dtc": row[1],
                "name": row[2],
                "confirmation_threshold": row[3],
                "healing_threshold": row[4],
            })

    return events


def _probe_debounce(wb) -> list[dict[str, Any]]:
    if "Debounce Configuration" not in wb.sheetnames:
        return []

    ws = wb["Debounce Configuration"]
    config = []

    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[0]:
            config.append({
                "event_id": row[0],
                "pre_passed": row[1],
                "pre_confirmed": row[2],
                "confirmed": row[3],
                "healed": row[4],
            })

    return config


def _probe_memory(wb) -> dict[str, Any]:
    if "Memory Allocation" not in wb.sheetnames:
        return {}

    ws = wb["Memory Allocation"]
    memory = {}

    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[0]:
            key = row[0].lower().replace(" ", "_")
            memory[key] = row[1]

    return memory


def _probe_snapshots(wb) -> list[dict[str, Any]]:
    if "Snapshot Record Class" not in wb.sheetnames:
        return []

    ws = wb["Snapshot Record Class"]
    snapshots = []

    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[0]:
            snapshots.append({
                "event_id": row[0],
                "snapshot_class": row[1],
                "data_elements": row[2] or "",
            })

    return snapshots


def _probe_edr(wb) -> list[dict[str, Any]]:
    if "Extended Data Record" not in wb.sheetnames:
        return []

    ws = wb["Extended Data Record"]
    edr = []

    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[0]:
            edr.append({
                "event_id": row[0],
                "edr_class": row[1],
                "data_elements": row[2] or "",
            })

    return edr


def _probe_nvm(wb) -> list[dict[str, Any]]:
    if "NvM Block Mapping" not in wb.sheetnames:
        return []

    ws = wb["NvM Block Mapping"]
    nvm = []

    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[0]:
            nvm.append({
                "event_id": row[0],
                "block_id": row[1],
                "block_type": row[2] or "",
            })

    return nvm
