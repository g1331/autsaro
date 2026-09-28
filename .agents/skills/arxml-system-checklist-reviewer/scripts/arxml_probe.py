"""ARXML system probe — extract data from xlsx for analysis."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from openpyxl import load_workbook


@dataclass
class ProbedARXML:
    """Probed state from ARXML system xlsx."""
    path: str = ""
    title: str = ""
    project: str = ""
    doc_id: str = ""
    revision: str = ""
    status: str = ""
    author: str = ""
    approver: str = ""

    sheet_names: list[str] = field(default_factory=list)
    has_table_of_contents: bool = False

    ecu_count: int = 0
    ecu_ids: set[str] = field(default_factory=set)

    signal_count: int = 0
    signal_ids: set[str] = field(default_factory=set)

    pdu_count: int = 0
    pdu_ids: set[str] = field(default_factory=set)
    signal_to_pdu_map: dict[str, str] = field(default_factory=dict)

    frame_count: int = 0
    frame_ids: set[str] = field(default_factory=set)
    pdu_to_frame_map: dict[str, str] = field(default_factory=dict)

    port_interface_count: int = 0
    port_interface_ids: set[str] = field(default_factory=set)

    composition_bindings: int = 0
    unbound_port_interfaces: int = 0

    mapping_count: int = 0
    ecu_pairs: set[tuple[str, str]] = field(default_factory=set)

    revision_history_rows: int = 0


def probe(xlsx_path: str) -> ProbedARXML:
    """Load and probe ARXML system workbook."""

    p = ProbedARXML(path=xlsx_path)

    try:
        wb = load_workbook(xlsx_path, data_only=True)
    except Exception as e:
        print(f"Error loading {xlsx_path}: {e}")
        return p

    p.sheet_names = wb.sheetnames

    # Title sheet
    if "Title" in wb.sheetnames:
        ws = wb["Title"]
        for row in ws.iter_rows(min_row=1, max_row=50, values_only=True):
            if row[0] and len(row) > 1:
                if "name" in str(row[0]).lower() and row[1]:
                    p.title = str(row[1])
                elif "project" in str(row[0]).lower() and row[1]:
                    p.project = str(row[1])
                elif "doc_id" in str(row[0]).lower() and row[1]:
                    p.doc_id = str(row[1])
                elif "revision" in str(row[0]).lower() and row[1]:
                    p.revision = str(row[1])
                elif "status" in str(row[0]).lower() and row[1]:
                    p.status = str(row[1])
                elif "author" in str(row[0]).lower() and row[1]:
                    p.author = str(row[1])
                elif "approver" in str(row[0]).lower() and row[1]:
                    p.approver = str(row[1])

    # Document Control sheet
    if "Document Control" in wb.sheetnames:
        ws = wb["Document Control"]
        p.revision_history_rows = sum(1 for row in ws.iter_rows(min_row=2, values_only=True) if row[0])

    # ECU Instances
    if "ECU Instances" in wb.sheetnames:
        ws = wb["ECU Instances"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0]:
                p.ecu_ids.add(str(row[0]))
        p.ecu_count = len(p.ecu_ids)

    # Signal Catalog
    if "Signal Catalog" in wb.sheetnames:
        ws = wb["Signal Catalog"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0]:
                p.signal_ids.add(str(row[0]))
        p.signal_count = len(p.signal_ids)

    # PDU Catalog
    if "PDU Catalog" in wb.sheetnames:
        ws = wb["PDU Catalog"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0]:
                p.pdu_ids.add(str(row[0]))
                if len(row) > 4 and row[4]:
                    signals = str(row[4]).split(", ")
                    for sig in signals:
                        sig_id = sig.strip()
                        if sig_id:
                            p.signal_to_pdu_map[sig_id] = str(row[0])
        p.pdu_count = len(p.pdu_ids)

    # Frame Catalog
    if "Frame Catalog" in wb.sheetnames:
        ws = wb["Frame Catalog"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0]:
                p.frame_ids.add(str(row[0]))
                if len(row) > 5 and row[5]:
                    pdus = str(row[5]).split(", ")
                    for pdu in pdus:
                        pdu_id = pdu.strip()
                        if pdu_id:
                            p.pdu_to_frame_map[pdu_id] = str(row[0])
        p.frame_count = len(p.frame_ids)

    # Port Interfaces
    if "Port Interfaces" in wb.sheetnames:
        ws = wb["Port Interfaces"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0]:
                p.port_interface_ids.add(str(row[0]))
        p.port_interface_count = len(p.port_interface_ids)

    # Composition SWC References
    if "Composition SWC References" in wb.sheetnames:
        ws = wb["Composition SWC References"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] and len(row) > 4:
                p.composition_bindings += 1
                if not row[4] or "unbound" in str(row[4]).lower():
                    p.unbound_port_interfaces += 1

    # System Mappings
    if "System Mappings" in wb.sheetnames:
        ws = wb["System Mappings"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0]:
                p.mapping_count += 1
                if len(row) > 2 and row[1] and row[2]:
                    src = str(row[1]).strip()
                    tgts = str(row[2]).split(", ")
                    for tgt in tgts:
                        tgt_id = tgt.strip()
                        if tgt_id:
                            p.ecu_pairs.add((src, tgt_id))

    p.has_table_of_contents = "System Header" in p.sheet_names

    return p
