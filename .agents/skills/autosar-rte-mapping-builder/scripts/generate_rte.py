"""
AUTOSAR RTE Mapping Builder — generates a multi-tab RTE mapping specification workbook.

Usage:
    python generate_rte.py <input.json> <output.xlsx>
"""

from __future__ import annotations

import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT_NAME = "Calibri"
NAVY = "1F3864"
LIGHT_BLUE = "D9E2F3"
THIN = Side(border_style="thin", color="BFBFBF")
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def style_title(ws, row, last_col, text):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(name=FONT_NAME, size=14, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 28

def style_header(ws, row, headers):
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_ALL
    ws.row_dimensions[row].height = 36

def build_title(wb, item):
    ws = wb.create_sheet("Title", 0)
    ws.sheet_view.showGridLines = False
    style_title(ws, 2, 4, "AUTOSAR RTE Mapping Specification Workbook")
    style_title(ws, 3, 4, item.get("name", "Unnamed RTE"))
    row = 5
    fields = [
        ("Project", item.get("project")),
        ("Document ID", item.get("doc_id")),
        ("Revision", item.get("revision", "1.0")),
        ("Date", item.get("date", str(date.today()))),
        ("Author", item.get("author")),
    ]
    for label, value in fields:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        row += 1

def build_document_control(wb):
    ws = wb.create_sheet("Document Control", 1)
    style_title(ws, 1, 4, "Document Control")
    headers = ["Change No.", "Revision", "Date", "Description"]
    style_header(ws, 2, headers)
    ws.cell(row=3, column=1, value=1)
    ws.cell(row=3, column=2, value="1.0")
    ws.cell(row=3, column=3, value=str(date.today()))

def build_ecu_integration_identity(wb, ecu):
    ws = wb.create_sheet("ECU Integration Identity", 2)
    style_title(ws, 1, 5, "ECU Integration Identity")
    row = 3
    fields = [
        ("ECU Name", ecu.get("name")),
        ("ECU ID", ecu.get("id")),
        ("Platform", ecu.get("platform")),
    ]
    for label, value in fields:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        row += 1

def build_data_mappings(wb, mappings):
    ws = wb.create_sheet("Data Mappings", 3)
    style_title(ws, 1, 7, "SWC Port to Com Signal Mappings")
    headers = ["SWC Name", "Port", "Port Type", "Com Signal", "Signal Type", "Length", "Byte Order"]
    style_header(ws, 2, headers)
    for idx, m in enumerate(mappings, start=3):
        ws.cell(row=idx, column=1, value=m.get("swc"))
        ws.cell(row=idx, column=2, value=m.get("port"))
        ws.cell(row=idx, column=3, value=m.get("port_type"))
        ws.cell(row=idx, column=4, value=m.get("com_signal"))
        ws.cell(row=idx, column=5, value=m.get("signal_type"))
        ws.cell(row=idx, column=6, value=m.get("length"))

def build_runnable_to_task(wb, runnables):
    ws = wb.create_sheet("Runnable-to-Task", 4)
    style_title(ws, 1, 7, "Runnable to OS Task Allocation")
    headers = ["SWC Name", "Runnable", "Allocated Task", "Trigger", "Execution Type", "Implicit/Explicit", "Notes"]
    style_header(ws, 2, headers)
    for idx, r in enumerate(runnables, start=3):
        ws.cell(row=idx, column=1, value=r.get("swc"))
        ws.cell(row=idx, column=2, value=r.get("runnable"))
        ws.cell(row=idx, column=3, value=r.get("task"))
        ws.cell(row=idx, column=4, value=r.get("trigger"))
        ws.cell(row=idx, column=5, value=r.get("execution_type"))

def build_schedulable_period(wb, tasks):
    ws = wb.create_sheet("Schedulable Period", 5)
    style_title(ws, 1, 6, "Schedulable Period and Priority")
    headers = ["Task Name", "Period (ms)", "Priority", "Runnable Count", "Total WCET (us)", "CPU Load %"]
    style_header(ws, 2, headers)
    for idx, t in enumerate(tasks, start=3):
        ws.cell(row=idx, column=1, value=t.get("name"))
        ws.cell(row=idx, column=2, value=t.get("period_ms"))
        ws.cell(row=idx, column=3, value=t.get("priority"))

def build_exclusive_area(wb, exclusive_areas):
    ws = wb.create_sheet("Exclusive Areas", 6)
    style_title(ws, 1, 6, "Exclusive Area to OS Resource Binding")
    headers = ["SWC Name", "Exclusive Area", "OS Resource", "Resource Type", "Protected By", "Notes"]
    style_header(ws, 2, headers)
    for idx, ea in enumerate(exclusive_areas, start=3):
        ws.cell(row=idx, column=1, value=ea.get("swc"))
        ws.cell(row=idx, column=2, value=ea.get("area"))
        ws.cell(row=idx, column=3, value=ea.get("os_resource"))
        ws.cell(row=idx, column=4, value=ea.get("resource_type"))

def build_mode_switch_mapping(wb, mode_switches):
    ws = wb.create_sheet("Mode-Switch Mapping", 7)
    style_title(ws, 1, 6, "Mode-Switch Mapping")
    headers = ["SWC Name", "Mode Port", "Mode Group", "Mapped Com Signal", "Transition Rules", "Notes"]
    style_header(ws, 2, headers)
    for idx, ms in enumerate(mode_switches, start=3):
        ws.cell(row=idx, column=1, value=ms.get("swc"))
        ws.cell(row=idx, column=2, value=ms.get("mode_port"))
        ws.cell(row=idx, column=3, value=ms.get("mode_group"))

def build_trigger_source_mappings(wb):
    ws = wb.create_sheet("Trigger Source Mappings", 8)
    style_title(ws, 1, 6, "Trigger Source Mappings")
    headers = ["Task Name", "Trigger Type", "Source Signal/Event", "Activation Condition", "Latency (ms)"]
    style_header(ws, 2, headers)

def build_mem_mapping(wb):
    ws = wb.create_sheet("Mem Mapping", 9)
    style_title(ws, 1, 5, "Memory Mapping (RAM Regions per Task)")
    headers = ["Task Name", "RAM Region", "Start Address", "Size (bytes)", "Accessibility"]
    style_header(ws, 2, headers)

def build_stack_sizing(wb):
    ws = wb.create_sheet("Stack Sizing", 10)
    style_title(ws, 1, 6, "Stack Sizing")
    headers = ["Task Name", "Allocated Stack (bytes)", "Measured Peak (bytes)", "Safety Margin %", "Recommendation"]
    style_header(ws, 2, headers)

def build_validation_rules(wb):
    ws = wb.create_sheet("Validation Rules", 11)
    style_title(ws, 1, 5, "Validation Rules")
    headers = ["Rule ID", "Rule Description", "Applies To", "Severity"]
    style_header(ws, 2, headers)

def build_references(wb):
    ws = wb.create_sheet("References", 12)
    style_title(ws, 1, 3, "References")
    row = 3
    refs = [
        ("AUTOSAR R22-11", "AUTOSAR Classic specification"),
        ("RTE Methodology", "scripts/references/rte_conventions.md"),
    ]
    for ref, desc in refs:
        ws.cell(row=row, column=1, value=ref).font = Font(bold=True)
        ws.cell(row=row, column=2, value=desc)
        row += 1

def main():
    if len(sys.argv) < 2:
        input_json = {
            "item": {"name": "Example RTE", "project": "Project X", "doc_id": "DOC-002"},
            "ecu": {"name": "ECU1", "id": "ECU_001", "platform": "AUTOSAR Classic"},
            "data_mappings": [],
            "runnable_to_task": [],
            "schedulable_tasks": [],
            "exclusive_areas": [],
            "mode_switches": [],
        }
    else:
        with open(sys.argv[1], encoding="utf-8") as f:
            input_json = json.load(f)

    wb = Workbook()
    wb.remove(wb.active)

    build_title(wb, input_json.get("item", {}))
    build_document_control(wb)
    build_ecu_integration_identity(wb, input_json.get("ecu", {}))
    build_data_mappings(wb, input_json.get("data_mappings", []))
    build_runnable_to_task(wb, input_json.get("runnable_to_task", []))
    build_schedulable_period(wb, input_json.get("schedulable_tasks", []))
    build_exclusive_area(wb, input_json.get("exclusive_areas", []))
    build_mode_switch_mapping(wb, input_json.get("mode_switches", []))
    build_trigger_source_mappings(wb)
    build_mem_mapping(wb)
    build_stack_sizing(wb)
    build_validation_rules(wb)
    build_references(wb)

    output_file = sys.argv[2] if len(sys.argv) > 2 else "rte_mapping.xlsx"
    _repository_notice(wb)
    wb.save(output_file)
    print(f"Generated {output_file}")



def _repository_notice(wb):
    import json
    from pathlib import Path
    import sys
    for sheet in wb:
        for row in sheet:
            for cell in row:
                if cell.value == "APPROVED":
                    cell.value = "DOCUMENT CHECKS COMPLETE - review required"
                elif cell.value == "CONDITIONAL APPROVAL":
                    cell.value = "DOCUMENT ISSUES - review required"
                elif cell.value in ("Internal review comments incorporated", "Released for architecture review", "Initial release"):
                    cell.value = "TEMPLATE PLACEHOLDER - not a recorded project event"
    if "03_Coding_Rules_Catalog" in wb.sheetnames:
        rules = wb["03_Coding_Rules_Catalog"]
        col = rules.max_column + 1
        rules.cell(1, col, "Repository validation status")
        for row in range(2, rules.max_row + 1):
            rules.cell(row, col, "UNVERIFIED UPSTREAM TEMPLATE - rule ID/text/severity/CWE require authoritative review")
    ws = wb.create_sheet("Repository Scope")
    ws.append(["Field", "Value"])
    ws.append(["Purpose", "Document generation/review only; NOT a runtime test or certification"])
    ws.append(["Evidence", "Separate CONTENT / STRUCTURE / DRAFT; missing evidence remains unassessed"])
    ws.append(["Version", "Upstream AUTOSAR R22-11 templates; repository R24-11 obligations need verification"])
    ws.append(["Thresholds", "Upstream thresholds are advisory, not repository acceptance gates"])
    ws.append(["Templates", "Fixed example/default content requires review; not project facts"])
    if len(sys.argv) > 1 and Path(sys.argv[1]).suffix.lower() == ".json":
        with open(sys.argv[1], encoding="utf-8") as stream:
            provenance = json.load(stream).get("provenance", {})
        for key, value in provenance.items():
            ws.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 100


if __name__ == "__main__":
    main()
